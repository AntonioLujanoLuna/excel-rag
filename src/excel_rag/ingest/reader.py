"""Read an ``.xlsx`` / ``.xlsm`` into raw sheets -- no macros, no link refresh, no evaluation.

Three rules govern this module:

* ``data_only=False`` gives the formula text, a second ``data_only=True`` load gives the last value
  Excel saved. **Both loads are needed** because openpyxl exposes them under one flag or the other,
  never together; the cached value is kept separately and labelled as cached, never recomputed.
* Macros are never loaded (``keep_vba=False``) and never executed -- openpyxl does not run VBA, and
  a workbook carrying a ``vbaProject.bin`` or a macro sheet is only *noticed*, so a formula pointing
  at a macro sheet resolves to an explicit ``macro_sheet`` gap rather than a fabricated edge.
* A declared dimension is not trusted. openpyxl already ignores ``<dimension ref="A1:XFD1048576">``,
  so nothing is allocated for it; we read the string from the package only to *report* that a
  hostile dimension was declared and ignored.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import openpyxl  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from .canonical import CellValue
from .errors import IngestError

_XLNS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RELNS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CTNS = "http://schemas.openxmlformats.org/package/2006/content-types"
_SHEET_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
_MACRO_REL_SUFFIX = "/xlMacrosheet"

#: A declared dimension larger than this many cells is reported as ignored, not honoured.
MAX_DECLARED_CELLS = 5_000_000
#: Refuse a package whose declared uncompressed size is absurd, before openpyxl reads it.
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024

_CURRENCY = re.compile(r"[$€£¥₹]")
_DIMENSION = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d{1,9})(?::\$?([A-Za-z]{1,3})\$?(\d{1,9}))?")


@dataclass(frozen=True, slots=True)
class RawTable:
    """An Excel table object (``ws.tables``) as read, before it is matched to a region."""

    name: str
    ref: str
    columns: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class RawDefinedName:
    """A defined name as read. ``local_sheet_id`` is the sheet scope, ``None`` for workbook-wide."""

    name: str
    attr_text: str | None
    local_sheet_id: int | None


@dataclass(frozen=True, slots=True)
class RawSheet:
    """A worksheet as read: populated cells, merges, tables, and package-level flags."""

    name: str
    visibility: str
    cells: dict[str, CellValue]
    merged_ranges: tuple[str, ...]
    tables: tuple[RawTable, ...]
    declared_dimension: str | None = None
    declared_dimension_flagged: bool = False
    is_macro_sheet: bool = False


@dataclass(frozen=True, slots=True)
class RawWorkbook:
    """Everything the reader extracted, before region detection."""

    path: str
    source_file: str
    source_sha256: str
    sheets: tuple[RawSheet, ...]
    defined_names: tuple[RawDefinedName, ...]
    macro_sheet_names: tuple[str, ...] = ()
    has_vba: bool = False
    warnings: tuple[str, ...] = field(default_factory=tuple)


@dataclass(frozen=True, slots=True)
class _SheetFact:
    name: str
    state: str
    part: str
    is_macro: bool


@dataclass(frozen=True, slots=True)
class _PackageFacts:
    sheets: tuple[_SheetFact, ...]
    declared_dimensions: dict[str, str]
    has_vba: bool

    @property
    def macro_names(self) -> frozenset[str]:
        return frozenset(sheet.name for sheet in self.sheets if sheet.is_macro)

    @property
    def by_name(self) -> dict[str, _SheetFact]:
        return {sheet.name: sheet for sheet in self.sheets}


def column_index(letters: str) -> int:
    """``A`` -> 1, ``Z`` -> 26, ``AA`` -> 27."""
    index = 0
    for char in letters.upper():
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def _dimension_area(ref: str) -> int:
    match = _DIMENSION.fullmatch(ref.strip())
    if match is None:
        return 0
    start_col, start_row, end_col, end_row = match.groups()
    min_col, min_row = column_index(start_col), int(start_row)
    max_col = column_index(end_col) if end_col else min_col
    max_row = int(end_row) if end_row else min_row
    if max_row < min_row or max_col < min_col:
        return 0
    return (max_row - min_row + 1) * (max_col - min_col + 1)


def merged_top_left(ref: str) -> str:
    """``"A1:B1"`` -> ``"A1"``: the only cell of a merged range that carries a value."""
    start = ref.split(":", 1)[0]
    return start.replace("$", "").upper()


def _read_package(data: bytes) -> _PackageFacts:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise IngestError(
            "not a valid .xlsx/.xlsm package (not a zip archive; the file may be corrupt, "
            "truncated, or password-protected)"
        ) from exc
    names = archive.namelist()
    uncompressed = sum(item.file_size for item in archive.infolist())
    if uncompressed > MAX_UNCOMPRESSED_BYTES:
        raise IngestError(
            f"refusing workbook: declared uncompressed size {uncompressed} bytes exceeds "
            f"the {MAX_UNCOMPRESSED_BYTES}-byte ceiling"
        )
    if "xl/workbook.xml" not in names:
        raise IngestError("not an Excel workbook package: xl/workbook.xml is missing")

    relationships: dict[str, tuple[str, str]] = {}
    if "xl/_rels/workbook.xml.rels" in names:
        rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for node in rels:
            resolved = node.attrib.get("Target", "")
            target = resolved.lstrip("/")
            if not resolved.startswith("/") and not target.startswith("xl/"):
                target = f"xl/{target}"
            relationships[node.attrib.get("Id", "")] = (node.attrib.get("Type", ""), target)

    content_types: dict[str, str] = {}
    if "[Content_Types].xml" in names:
        types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        for node in types:
            part = node.attrib.get("PartName")
            if part:
                content_types[part.lstrip("/")] = node.attrib.get("ContentType", "")

    sheets: list[_SheetFact] = []
    declared: dict[str, str] = {}
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    sheet_container = workbook.find(f"{{{_XLNS}}}sheets")
    if sheet_container is not None:
        for node in sheet_container:
            name = node.attrib.get("name", "")
            rel_id = node.attrib.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", ""
            )
            rel_type, target = relationships.get(rel_id, ("", ""))
            is_macro = rel_type.endswith(_MACRO_REL_SUFFIX) or content_types.get(
                target, ""
            ).endswith("macrosheet+xml")
            sheets.append(
                _SheetFact(
                    name=name,
                    state=node.attrib.get("state", "visible"),
                    part=target,
                    is_macro=is_macro,
                )
            )
            if target in names:
                try:
                    part_root = ElementTree.fromstring(archive.read(target))
                except ElementTree.ParseError:
                    part_root = None
                if part_root is not None:
                    dimension = part_root.find(f"{{{_XLNS}}}dimension")
                    if dimension is not None and dimension.attrib.get("ref"):
                        declared[target] = dimension.attrib["ref"]

    has_vba = any(name.endswith("vbaProject.bin") for name in names)
    return _PackageFacts(sheets=tuple(sheets), declared_dimensions=declared, has_vba=has_vba)


def _load(data: bytes, *, data_only: bool) -> Any:
    try:
        return openpyxl.load_workbook(
            io.BytesIO(data), data_only=data_only, keep_vba=False, read_only=False
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise IngestError(f"could not read the workbook: {exc}") from exc


def read_workbook(path: str | Path) -> RawWorkbook:
    """Read ``path`` into :class:`RawWorkbook`, or raise :class:`IngestError`.

    Never loads or executes macros, never refreshes external links, never evaluates a formula.
    """
    source = Path(path)
    if not source.is_file():
        raise IngestError(f"no such workbook: {source}")
    data = source.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    facts = _read_package(data)

    formula_wb = _load(data, data_only=False)
    try:
        cached_wb = _load(data, data_only=True)
    except IngestError:
        cached_wb = None

    macro_names = facts.macro_names
    sheets: list[RawSheet] = []
    for ws in formula_wb.worksheets:
        name = ws.title
        fact = facts.by_name.get(name)
        declared_ref = facts.declared_dimensions.get(fact.part) if fact else None
        if name in macro_names:
            sheets.append(
                RawSheet(
                    name=name,
                    visibility=fact.state if fact else ws.sheet_state,
                    cells={},
                    merged_ranges=(),
                    tables=(),
                    declared_dimension=declared_ref,
                    declared_dimension_flagged=False,
                    is_macro_sheet=True,
                )
            )
            continue

        merged_ranges = tuple(str(item) for item in ws.merged_cells.ranges)
        master_ranges = {merged_top_left(item): item for item in merged_ranges}
        cached_ws = (
            cached_wb[name] if cached_wb is not None and name in cached_wb.sheetnames else None
        )
        cells: dict[str, CellValue] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                data_type = cell.data_type or "n"
                formula = str(cell.value) if data_type == "f" else None
                cached_value: Any = None
                if formula is not None and cached_ws is not None:
                    cached_value = cached_ws[cell.coordinate].value
                number_format = cell.number_format or "General"
                cells[cell.coordinate] = CellValue(
                    sheet_name=name,
                    coordinate=cell.coordinate,
                    row=cell.row,
                    column=cell.column,
                    value=cell.value if formula is None else None,
                    formula=formula,
                    cached_value=cached_value,
                    data_type=data_type,
                    number_format=number_format,
                    is_date=bool(cell.is_date),
                    is_percentage="%" in number_format,
                    is_currency=bool(_CURRENCY.search(number_format)),
                    merged_range=master_ranges.get(cell.coordinate),
                )
        tables: list[RawTable] = []
        for table_name in list(ws.tables):
            table = ws.tables[table_name]
            columns = tuple(column.name or "" for column in (table.tableColumns or ()))
            tables.append(
                RawTable(
                    name=table.displayName or table_name,
                    ref=table.ref or "",
                    columns=columns,
                )
            )

        flagged = (
            declared_ref is not None
            and _dimension_area(declared_ref) > MAX_DECLARED_CELLS
            and _dimension_area(declared_ref) > 4 * max(1, len(cells))
        )
        sheets.append(
            RawSheet(
                name=name,
                visibility=fact.state if fact else ws.sheet_state,
                cells=cells,
                merged_ranges=merged_ranges,
                tables=tuple(tables),
                declared_dimension=declared_ref,
                declared_dimension_flagged=flagged,
                is_macro_sheet=False,
            )
        )

    defined_names = tuple(
        RawDefinedName(name=name, attr_text=dn.attr_text, local_sheet_id=dn.localSheetId)
        for name, dn in formula_wb.defined_names.items()
    )

    return RawWorkbook(
        path=str(source),
        source_file=source.name,
        source_sha256=digest,
        sheets=tuple(sheets),
        defined_names=defined_names,
        macro_sheet_names=tuple(sorted(macro_names)),
        has_vba=facts.has_vba,
        warnings=tuple(
            f"sheet {sheet.name!r} declared dimension {sheet.declared_dimension!r}; ignored"
            for sheet in sheets
            if sheet.declared_dimension_flagged
        ),
    )
