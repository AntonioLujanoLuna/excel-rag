"""Formula evaluation on request: saved values where no change reaches, recomputed values where
one does, and unknown -- with the reason -- where a value cannot be computed."""

from __future__ import annotations

import random
import sys
from pathlib import Path

import pytest

pytest.importorskip("formulas")

sys.path.insert(0, str(Path(__file__).parent / "fixtures"))

import make_fixtures as mk

from excel_rag.cli import main as cli_main
from excel_rag.context import CALCULATE_DEFINITION, WorkbookSession
from excel_rag.evaluate.sample import build_sample_workbook
from excel_rag.workbook import load_workbook
from excel_rag.workbook.arithmetic import ExcelError, Fallback, compile_arithmetic
from excel_rag.workbook.calc import (
    CalcInputError,
    CalcUnavailable,
    CellResult,
    calculate,
    check_saved_values,
    format_calculation,
    format_check,
)

MODEL = "'Model Sheet'"


@pytest.fixture
def model(tmp_path: Path):
    return load_workbook(mk.calc_model(tmp_path))


def _by_cell(results: tuple[CellResult, ...]) -> dict[str, CellResult]:
    return {result.coordinate: result for result in results}


class TestSavedValues:
    def test_with_no_change_every_formula_is_excels_saved_value(self, model) -> None:
        calculation = calculate(model, f"{MODEL}!A1:A12")
        assert calculation.recalculated == 0
        assert {result.status for result in calculation.results} == {"saved"}
        assert _by_cell(calculation.results)["A2"].value == 44

    def test_a_workbook_saved_without_values_is_computed(self, tmp_path: Path) -> None:
        planning = load_workbook(build_sample_workbook(tmp_path))
        result = calculate(planning, "Forecast!B4").results[0]
        # SUM(Actuals!C2:C25) is 31,800; /2 * (1 + 5%)^1.
        assert (result.status, result.value) == ("recalculated", 16695)
        grown = calculate(planning, "Forecast!B4", {"Assumptions!B4": "7%"}).results[0]
        assert grown.value == pytest.approx(17013)


class TestWhatIf:
    def test_a_change_moves_what_reads_it_and_nothing_else(self, model) -> None:
        results = _by_cell(calculate(model, f"{MODEL}!A1:A12", {"Inputs!B2": 20}).results)
        assert (results["A1"].value, results["A1"].status, results["A1"].saved) == (
            80,
            "recalculated",
            40,
        )
        assert results["A2"].value == 88  # through the defined name Rate
        assert results["A4"].value == 160  # LET, as Excel stores it (_xlfn.LET, _xlpm.x)
        assert results["A3"].status == "saved"  # YEAR(Inputs!B4) is not reached
        assert results["A12"].status == "saved"

    def test_a_circular_pair_no_change_feeds_keeps_its_saved_values(self, model) -> None:
        results = _by_cell(calculate(model, f"{MODEL}!A7:A8", {"Inputs!B2": 20}).results)
        assert [(results[c].value, results[c].status) for c in ("A7", "A8")] == [
            (1, "saved"),
            (2, "saved"),
        ]

    def test_a_structured_reference_is_computed_over_its_column(self, model) -> None:
        result = calculate(model, f"{MODEL}!A10", {"Inputs!E3": 25}).results[0]
        assert (result.value, result.status) == (65, "recalculated")

    def test_several_changes_at_once(self, model) -> None:
        changes = {"Inputs!B2": "20", "Inputs!B1": "15%"}
        assert calculate(model, f"{MODEL}!A2", changes).results[0].value == pytest.approx(92)

    def test_an_error_value_propagates(self, model) -> None:
        results = _by_cell(calculate(model, f"{MODEL}!A1:A2", {"Inputs!B3": "#N/A"}).results)
        assert results["A1"].value == "#N/A" and results["A2"].value == "#N/A"

    def test_a_long_chain_is_not_a_recursion_limit(self, tmp_path: Path) -> None:
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Ledger"
        sheet["C1"] = 1
        sheet["B1"] = "=C1"
        for row in range(2, 3001):
            sheet[f"A{row}"] = 1
            sheet[f"B{row}"] = f"=B{row - 1}+A{row}"
        path = tmp_path / "ledger.xlsx"
        workbook.save(path)
        result = calculate(load_workbook(path), "Ledger!B3000", {"Ledger!C1": 10}).results[0]
        assert result.value == 3009


class TestUnknown:
    def test_what_cannot_be_computed_is_unknown_with_the_reason(self, model) -> None:
        results = _by_cell(calculate(model, f"{MODEL}!A1:A12", {"Inputs!B2": 20}).results)
        for coordinate, words in (("A5", "INDIRECT"), ("A9", "OFFSET"), ("A11", "user-defined")):
            result = results[coordinate]
            assert result.status == "unknown" and result.value is None
            assert words in (result.reason or "")
        # What Excel last saved is still shown beside it.
        assert results["A5"].saved == 11

    def test_a_circular_reference_is_reported_not_iterated(self, model) -> None:
        results = _by_cell(calculate(model, f"{MODEL}!A7:A8", recalculate_all=True).results)
        assert all(result.status == "unknown" for result in results.values())
        assert "circular reference" in (results["A8"].reason or "")

    def test_an_xlsb_formula_with_no_text_is_saved_until_a_change_could_reach_it(
        self, tmp_path: Path
    ) -> None:
        pytest.importorskip("pyxlsb")
        book = load_workbook(mk.xlsb_workbook(tmp_path))
        assert calculate(book, "Sales!C2").results[0].status == "saved"
        changed = calculate(book, "Sales!C2", {"Sales!B2": 1}).results[0]
        assert changed.status == "unknown" and "not stored" in (changed.reason or "")


class TestArrays:
    def test_an_array_formula_and_its_spill_references_are_recomputed(self, tmp_path: Path) -> None:
        book = load_workbook(mk.array_formula(tmp_path))
        results = _by_cell(calculate(book, "Calc!A2:E4", {"Calc!A3": 5}).results)
        assert results["B3"].value == 10 and results["B3"].status == "recalculated"
        assert results["B4"].value == 6
        assert results["C2"].value == 18  # SUM(_xlfn.ANCHORARRAY(B2))
        assert results["D2"].value == 18  # SUM(B2#)
        # A what-if data table reruns the model with substituted inputs: not computed here.
        assert results["E2"].status == "unknown"


class TestCheck:
    def test_a_full_recalculation_agrees_with_excels_saved_values(self, model) -> None:
        check = check_saved_values(model)
        assert check.mismatches == ()
        assert "'Model Sheet'!A6" in check.volatile  # NOW() differs by design, not by error
        text = format_check(check)
        assert "0 differ" in text and "1 volatile" in text

    def test_a_wrong_saved_value_is_a_mismatch(self, tmp_path: Path) -> None:
        book = load_workbook(mk.cached_value(tmp_path))
        check = check_saved_values(book)
        assert [result.qualified for result in check.mismatches] == ["Calc!C1"]
        assert "Excel saved 999" in format_check(check)


class TestInputs:
    @pytest.mark.parametrize(
        ("typed", "expected"),
        [("7%", 0.07), ("1,200", 1200), ("TRUE", True), ("-3.5", -3.5), ("north", "north")],
    )
    def test_a_change_is_taken_as_typed(self, model, typed: str, expected: object) -> None:
        calculation = calculate(model, "Inputs!B2", {"Inputs!B2": typed})
        assert calculation.results[0].value == expected
        assert calculation.results[0].status == "changed"

    def test_clearing_a_cell(self, model) -> None:
        result = calculate(model, f"{MODEL}!A1", {"Inputs!B2": ""}).results[0]
        assert result.value == 0

    @pytest.mark.parametrize(
        ("targets", "changes", "message"),
        [
            ("Nope!A1", {}, "no sheet named 'Nope'"),
            ("A1", {}, "with its sheet"),
            (f"{MODEL}!A1", {"Inputs!B2": "=B3*2"}, "not a formula"),
            (f"{MODEL}!A1", {"Inputs!B2:B3": "1"}, "one cell"),
        ],
    )
    def test_a_bad_target_or_change_is_refused(
        self, model, targets: str, changes: dict[str, str], message: str
    ) -> None:
        with pytest.raises(CalcInputError, match=message):
            calculate(model, targets, changes)

    def test_without_the_extra_it_says_how_to_install_it(self, model, monkeypatch) -> None:
        monkeypatch.setitem(sys.modules, "formulas", None)
        with pytest.raises(CalcUnavailable, match="calc extra"):
            calculate(model, f"{MODEL}!A1")


class TestLimits:
    def test_an_evaluation_limit_stops_and_says_so(self, model) -> None:
        calculation = calculate(model, f"{MODEL}!A1:A4", {"Inputs!B2": 20}, max_evaluations=1)
        assert calculation.stopped is not None and "evaluation limit" in calculation.stopped
        statuses = [result.status for result in calculation.results]
        assert "unknown" in statuses and "recalculated" in statuses
        assert "Stopped early" in format_calculation(calculation)

    def test_a_time_limit_stops_and_says_so(self, model) -> None:
        calculation = calculate(model, f"{MODEL}!A1", {"Inputs!B2": 20}, timeout_seconds=0)
        assert calculation.stopped == "the time limit was reached"
        assert calculation.results[0].status == "unknown"


class TestArithmeticFastPath:
    @pytest.mark.parametrize(
        ("template", "values", "expected"),
        [
            ("=-_REF0^2", (2,), 4.0),  # negation binds tighter than ^, as in Excel
            ("=_REF0^3^2", (2,), 64.0),  # ^ is left-associative in Excel
            ("=_REF0%^2", (50,), 0.25),
            ("=-_REF0%", (5,), -0.05),
            ("=(_REF0 - 1.5E+3)/_REF1", (2000, 4), 125.0),
            ("=_REF0+_REF1", (True, None), None),  # None stands for an empty cell below
        ],
    )
    def test_excel_precedence(self, template, values, expected) -> None:
        import schedula

        arguments = [schedula.EMPTY if value is None else value for value in values]
        function = compile_arithmetic(template, [f"_REF{i}" for i in range(len(values))])
        assert function is not None
        assert function(*arguments) == (1.0 if expected is None else expected)

    @pytest.mark.parametrize(
        ("template", "values", "code"),
        [
            ("=_REF0^_REF1", (0, 0), "#NUM!"),
            ("=_REF0^(1/3)", (-8,), "#NUM!"),
            ("=_REF0^-1", (0,), "#DIV/0!"),
            ("=_REF0/0+_REF1", (1, "#N/A"), "#DIV/0!"),  # the first error met wins
            ("=_REF0*1E+300*1E+300", (1,), "#NUM!"),
        ],
    )
    def test_excel_errors(self, template, values, code) -> None:
        from formulas.tokens.operand import XlError

        arguments = [XlError(value) if str(value).startswith("#") else value for value in values]
        function = compile_arithmetic(template, [f"_REF{i}" for i in range(len(values))])
        assert function is not None
        with pytest.raises(ExcelError) as raised:
            function(*arguments)
        assert raised.value.code == code

    def test_text_is_left_to_the_general_evaluator(self) -> None:
        function = compile_arithmetic("=_REF0+1", ["_REF0"])
        assert function is not None
        with pytest.raises(Fallback):
            function("3")

    @pytest.mark.parametrize(
        "template", ["=SUM(_REF0)", '=_REF0&"x"', "=_REF0>1", "=_REF0 _REF1", "=(_REF0"]
    )
    def test_more_than_arithmetic_is_not_compiled(self, template) -> None:
        assert compile_arithmetic(template, ["_REF0", "_REF1"]) is None

    def test_it_agrees_with_the_general_evaluator(self) -> None:
        """Random arithmetic, where both define the result, gives the same number both ways."""
        import formulas

        generator = random.Random(7)

        def expression(depth: int) -> str:
            if depth == 0 or generator.random() < 0.3:
                return generator.choice(["_REF0", "_REF1", "_REF2", "2", "0.5", "3"])
            operator = generator.choice(["+", "-", "*", "/"])
            left, right = expression(depth - 1), expression(depth - 1)
            text = f"{left}{operator}{right}"
            return f"({text})" if generator.random() < 0.5 else text

        compared = 0
        for _ in range(300):
            template = "=" + expression(4)
            names = [name for name in ("_REF0", "_REF1", "_REF2") if name in template]
            values = [generator.choice([1.5, -2, 4, 7.25]) for _ in names]
            fast = compile_arithmetic(template, names)
            assert fast is not None, template
            general = formulas.Parser().ast(template)[1].compile()
            ordered = [values[names.index(str(name))] for name in general.inputs]
            try:
                expected = float(general(*ordered))
            except (TypeError, ValueError):
                continue  # an error value: compared in test_excel_errors
            try:
                actual = fast(*values)
            except ExcelError:
                continue
            assert actual == pytest.approx(expected, rel=1e-12), template
            compared += 1
        assert compared > 200


class TestTool:
    def test_the_definition_is_strict(self) -> None:
        schema = CALCULATE_DEFINITION["input_schema"]
        assert CALCULATE_DEFINITION["strict"] is True
        assert schema["required"] == ["sheet", "range", "changes"]
        item = schema["properties"]["changes"]["items"]
        assert item["additionalProperties"] is False and item["required"] == ["cell", "value"]

    def test_a_what_if_through_the_tool(self, tmp_path: Path) -> None:
        session = WorkbookSession.load(mk.calc_model(tmp_path))
        outcome = session.call(
            "calculate",
            {
                "sheet": "Model Sheet",
                "range": "A1:A2",
                "changes": [{"cell": "Inputs!B2", "value": "20"}],
            },
        )
        assert not outcome.is_error
        assert "With Inputs!B2 = 20:" in outcome.content
        assert "'Model Sheet'!A2: 88 (recalculated; Excel saved 44)" in outcome.content

    @pytest.mark.parametrize(
        "changes",
        [None, "Inputs!B2=20", [{"cell": "Inputs!B2"}], [{"cell": "", "value": "1"}]],
    )
    def test_a_malformed_change_is_a_correctable_error(self, tmp_path: Path, changes) -> None:
        session = WorkbookSession.load(mk.calc_model(tmp_path))
        outcome = session.call(
            "calculate", {"sheet": "Model Sheet", "range": "A1", "changes": changes}
        )
        assert outcome.is_error

    def test_a_change_to_an_unknown_sheet_is_a_correctable_error(self, tmp_path: Path) -> None:
        session = WorkbookSession.load(mk.calc_model(tmp_path))
        outcome = session.call(
            "calculate",
            {"sheet": "Inputs", "range": "B2", "changes": [{"cell": "Nope!A1", "value": "1"}]},
        )
        assert outcome.is_error and "no sheet named 'Nope'" in outcome.content

    def test_offered_only_when_asked_or_installed(self) -> None:
        names = [tool["name"] for tool in WorkbookSession.tool_definitions(calculate=False)]
        assert "calculate" not in names
        assert WorkbookSession.tool_definitions(calculate=True)[-1]["name"] == "calculate"


class TestCommandLine:
    def test_calc_with_a_change(self, tmp_path: Path, capsys) -> None:
        path = str(mk.calc_model(tmp_path))
        assert cli_main(["calc", path, f"{MODEL}!A1", "--set", "Inputs!B2=20"]) == 0
        assert "'Model Sheet'!A1: 80 (recalculated; Excel saved 40)" in capsys.readouterr().out

    def test_check_exits_3_on_a_mismatch(self, tmp_path: Path, capsys) -> None:
        assert cli_main(["calc", str(mk.calc_model(tmp_path)), "--check"]) == 0
        assert cli_main(["calc", str(mk.cached_value(tmp_path)), "--check"]) == 3
        assert "1 differ" in capsys.readouterr().out

    def test_json(self, tmp_path: Path, capsys) -> None:
        import json

        assert cli_main(["calc", str(mk.calc_model(tmp_path)), f"{MODEL}!A1", "--json"]) == 0
        payload = json.loads(capsys.readouterr().out)
        assert payload["results"][0]["status"] == "saved"

    @pytest.mark.parametrize(
        ("argv", "code"),
        [(["--set", "no-equals"], 2), ([], 2), (["Nope!A1"], 2)],
    )
    def test_bad_arguments(self, tmp_path: Path, argv: list[str], code: int) -> None:
        assert cli_main(["calc", str(mk.calc_model(tmp_path)), *argv]) == code
