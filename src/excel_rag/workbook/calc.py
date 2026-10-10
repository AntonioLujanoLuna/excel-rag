"""Evaluate a workbook's formulas on request: what a cell is now, or would be if inputs changed.

Everything else in this package reads what Excel *saved*. This module computes, and only when a
caller asks (``excel-rag calc``, the ``calculate`` tool): given changes to some cells, which values
move and to what. Its results are labelled as computed by excel-rag, never passed off as Excel's.

How it stays honest and safe:

* **Excel's own values wherever they still hold.** A formula cell is recomputed only when a change
  can reach it (or it has no saved value, or a full recalculation was asked for); every other cell
  keeps the value Excel saved. Reachability is decided per cell from what each formula reads, so a
  circular pair no change feeds keeps its saved values too.
* **One formula at a time, never the file.** Each formula is compiled by the ``formulas`` library
  from its text -- which builds a function graph, it does not execute generated Python -- after its
  references are rewritten to plain A1 rectangles by this package's own resolver (structured table
  references, defined names, spill references). Cell values come from the reader, so external
  links are never followed and macros never run.
* **Unknown is a result, not a guess.** A formula whose precedents depend on runtime strings
  (``INDIRECT``, ``OFFSET``), on another workbook, a user-defined function, a 3-D reference, or a
  function the evaluator lacks cannot be recomputed. If a change can reach it, its value is
  ``unknown`` with the reason, and so is everything computed from it. Circular references are
  reported, not iterated.
* **Bounded.** Formula evaluations, the cells one range may supply and the wall-clock are capped;
  a stop is reported as such.

Install the ``calc`` extra (``formulas``) to use it.
"""

from __future__ import annotations

import math
import re
import time
import warnings
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from datetime import time as dtime
from typing import Any, Literal

from openpyxl.formula.tokenizer import (  # type: ignore[import-untyped]
    Token,
    Tokenizer,
    TokenizerError,
)

from ..models import A1Range, UnresolvedReason
from .arithmetic import ExcelError, Fallback, compile_arithmetic
from .canonical import CellValue, FormulaEntry, SheetModel, WorkbookModel, column_letter
from .errors import WorkbookError
from .formulas import (
    _SPILL_ANCHOR_RE,
    _SPILL_RE,
    _STRING_RE,
    FormulaContext,
    _blank,
    _bound_names,
    _function_gaps,
    parse_formula,
)

#: Formula evaluations one calculation may perform.
MAX_EVALUATIONS = 100_000
#: Cells one range input may supply to a formula.
MAX_RANGE_CELLS = 1_000_000
#: Wall-clock budget for one calculation.
TIMEOUT_SECONDS = 30.0

#: Functions whose value changes on every calculation: a recomputed result is as of now.
VOLATILE_FUNCTIONS = frozenset({"NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY"})

#: Excel's error values, as a cached cell or a change may hold them.
ERROR_VALUES = frozenset(
    {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
)

#: Why a formula's static edges are not the cells it reads, so it cannot be recomputed.
_BLOCKING_REASONS = frozenset(
    {
        UnresolvedReason.INDIRECT,
        UnresolvedReason.VOLATILE_OFFSET,
        UnresolvedReason.EXTERNAL_LINK,
        UnresolvedReason.MACRO_SHEET,
        UnresolvedReason.UNSUPPORTED_FUNCTION,
        UnresolvedReason.DYNAMIC_ARRAY,
        UnresolvedReason.OUT_OF_RANGE,
        UnresolvedReason.MALFORMED,
    }
)

Status = Literal["input", "changed", "saved", "recalculated", "unknown"]


class CalcUnavailable(WorkbookError):
    """The ``calc`` extra (the ``formulas`` library) is not installed."""


class CalcInputError(ValueError):
    """A change or target the caller wrote that does not name a cell of the workbook."""


@dataclass(frozen=True, slots=True)
class CellResult:
    """One cell after a calculation.

    ``status`` says where ``value`` came from: ``input`` (a value cell, unchanged), ``changed`` (a
    cell the caller set), ``saved`` (a formula Excel computed, which no change reaches),
    ``recalculated`` (computed here) or ``unknown`` (``reason`` says why). ``saved`` is Excel's
    last saved value for a formula cell, for comparison.
    """

    sheet: str
    coordinate: str
    value: Any
    status: Status
    saved: Any = None
    formula: str | None = None
    reason: str | None = None

    @property
    def qualified(self) -> str:
        return _qualified(self.sheet, self.coordinate)


@dataclass(frozen=True)
class Calculation:
    """What one calculation produced: the target cells, and how it went."""

    results: tuple[CellResult, ...]
    changes: tuple[tuple[str, str, Any], ...]
    recalculated: int
    #: Cells whose value changes on every calculation (``NOW``...), among those recomputed.
    volatile: tuple[str, ...] = ()
    #: Why the calculation stopped early, when it did; the cells it did not reach are unknown.
    stopped: str | None = None
    #: Formula cells whose recomputed value differs from Excel's saved one (``recalculate_all``).
    mismatches: tuple[CellResult, ...] = field(default_factory=tuple)


@dataclass(frozen=True, slots=True)
class _Unknown:
    reason: str


@dataclass(slots=True)
class _Plan:
    """How to compute one formula cell.

    ``inputs`` are the rectangles it reads: what the compiled ``function`` is fed or -- for a
    formula the evaluator cannot run, ``blocked`` saying why -- its static edges, which still say
    whether a change reaches it. ``unbounded`` marks a formula that may read cells no edge names
    (``INDIRECT``, ``OFFSET``, a user-defined function, a data table): any change may reach it.
    """

    function: Any
    inputs: tuple[A1Range, ...]
    volatile: bool = False
    blocked: str | None = None
    unbounded: bool = False
    #: The arithmetic fast path for this formula, when it is only arithmetic over cells.
    fast: Any = None


def calculate(
    model: WorkbookModel,
    targets: str | Sequence[str],
    changes: Mapping[str, Any] | None = None,
    *,
    recalculate_all: bool = False,
    max_evaluations: int = MAX_EVALUATIONS,
    max_range_cells: int = MAX_RANGE_CELLS,
    timeout_seconds: float = TIMEOUT_SECONDS,
) -> Calculation:
    """Compute ``targets`` (``"Sheet!A1"``, ``"Sheet!B2:D9"``, or a list) under ``changes``.

    ``changes`` maps a cell (``"Assumptions!B4"``) to its new value: a number, text, a boolean, an
    Excel error (``"#N/A"``) or ``None`` to clear it. A string that reads as a number or a boolean
    is taken as one, as Excel takes typed input; a leading ``=`` is refused (a change sets a
    value, not a formula). With ``recalculate_all`` every formula is recomputed rather than read
    from Excel's saved values, and differences from them are reported as ``mismatches``.

    Raises :class:`CalcUnavailable` without the ``calc`` extra and :class:`CalcInputError` for a
    target or change that names no cell of the workbook.
    """
    engine = _Engine(
        model,
        _parse_changes(model, changes or {}),
        recalculate_all=recalculate_all,
        max_evaluations=max_evaluations,
        max_range_cells=max_range_cells,
        deadline=time.monotonic() + timeout_seconds,
    )
    rectangles = [_parse_target(model, text) for text in _as_list(targets)]
    results = [engine.result(sheet, coordinate) for sheet, coordinate in _cells(model, rectangles)]
    mismatches: list[CellResult] = []
    if recalculate_all:
        # A volatile cell (NOW...) differs from its saved value by design, not by error.
        mismatches = [
            result
            for result in results
            if _differs(result) and result.qualified not in engine.volatile
        ]
    return Calculation(
        results=tuple(results),
        changes=tuple(
            (sheet, coordinate, _plain(value))
            for (sheet, coordinate), value in engine.changes.items()
        ),
        recalculated=engine.evaluations,
        volatile=tuple(sorted(engine.volatile)),
        stopped=engine.stopped,
        mismatches=tuple(mismatches),
    )


def check_saved_values(model: WorkbookModel, **limits: Any) -> Calculation:
    """Recompute every formula and compare with what Excel saved: how far excel-rag's evaluation
    can be trusted on this workbook. Cells with no saved value are not compared."""
    targets = [
        _qualified(sheet.name, coordinate)
        for sheet in model.sheets
        if not sheet.is_macro_sheet
        for coordinate, cell in sorted(sheet.cells.items(), key=lambda item: _order(item[1]))
        if cell.formula is not None
    ]
    if not targets:
        return Calculation(results=(), changes=(), recalculated=0)
    return calculate(model, targets, recalculate_all=True, **limits)


# -------------------------------------------------------------------------------------------------
# The engine: decide which cells a change reaches, then compute only those
# -------------------------------------------------------------------------------------------------
Key = tuple[str, str]


class _Engine:
    def __init__(
        self,
        model: WorkbookModel,
        changes: dict[Key, Any],
        *,
        recalculate_all: bool,
        max_evaluations: int,
        max_range_cells: int,
        deadline: float,
    ) -> None:
        self.library = _load_formulas()
        self.model = model
        self.changes = changes
        self.recalculate_all = recalculate_all
        self.max_evaluations = max_evaluations
        self.max_range_cells = max_range_cells
        self.deadline = deadline
        self.sheets: dict[str, SheetModel] = {sheet.name: sheet for sheet in model.sheets}
        self.entries: dict[Key, FormulaEntry] = {
            (sheet.name, coordinate): entry
            for sheet in model.sheets
            for entry in sheet.formulas
            for coordinate in entry.member_coordinates
        }
        #: Formula entries a change might reach: a cheap superset that rules most cells out.
        self.reachable = _reachable_entries(model, changes)
        #: Whether a change reaches each computed cell (decided exactly, on demand).
        self.dirty: dict[Key, bool] = {}
        self.plans: dict[Key, _Plan] = {}
        self.compiled: dict[str, Any] = {}
        self.fast: dict[str, Any] = {}
        self.operands: Operands = {}
        self.memo: dict[Key, Any] = {}
        self.status: dict[Key, Status] = {}
        #: An array formula master's whole recomputed result, for the cells of its extent.
        self.arrays: dict[Key, Any] = {}
        self.evaluations = 0
        self.volatile: set[str] = set()
        self.stopped: str | None = None

    # -- results ------------------------------------------------------------------------------
    def result(self, sheet: str, coordinate: str) -> CellResult:
        key = (sheet, coordinate)
        value = self.value(key)
        cell = self._cell(key)
        formula = cell.formula if cell is not None else None
        saved = _plain(cell.cached_value) if cell is not None and cell.computed else None
        if isinstance(value, _Unknown):
            return CellResult(sheet, coordinate, None, "unknown", saved, formula, value.reason)
        return CellResult(
            sheet, coordinate, _plain(value), self.status.get(key, "input"), saved, formula
        )

    # -- does a change reach this cell? ---------------------------------------------------------
    def is_dirty(self, key: Key) -> bool:
        if key in self.changes:
            return True
        if key in self.dirty:
            return self.dirty[key]
        quick = self._quick_dirty(key)
        if quick is not None:
            self.dirty[key] = quick
            return quick
        self._mark(key)
        return self.dirty[key]

    def _quick_dirty(self, key: Key) -> bool | None:
        """Whether a change reaches ``key`` when that needs no look at its precedents."""
        if key in self.changes:
            return True
        cell = self._cell(key)
        if cell is None or not cell.computed:
            return False
        if cell.array_master is not None:
            return None  # as its master is
        if cell.formula is None:  # an .xlsb formula: no text, so no telling what it reads
            return self.recalculate_all or bool(self.changes)
        if self.recalculate_all or cell.cached_value is None:
            return True
        entry = self.entries.get(key)
        if entry is None or entry.node_id not in self.reachable:
            return False
        return None

    def _mark(self, start: Key) -> None:
        """Decide reachability for ``start`` and the undecided cells upstream of it, exactly: a
        cell is reached when a change is among the cells it reads, or a reached cell is. A cycle
        is reached only if something feeding it is."""
        cone: dict[Key, list[Key]] = {}
        seeds: set[Key] = set()
        stack = [start]
        while stack:
            key = stack.pop()
            if key in cone or key in self.dirty:
                continue
            if len(cone) % 1024 == 0 and self._out_of_time():
                # What is left undecided counts as reached, and so is reported not computed.
                cone[key] = []
                seeds.update([key, *stack])
                cone.update({pending: [] for pending in stack if pending not in cone})
                break
            plan = self._plan(key)
            reads = self._read_cells(plan.inputs, include_changes=True)
            cone[key] = reads
            if plan.unbounded and self.changes:
                seeds.add(key)
            for dep in reads:
                if dep in cone or dep in self.dirty:
                    continue
                quick = self._quick_dirty(dep)
                if quick is None:
                    stack.append(dep)
                else:
                    self.dirty[dep] = quick
        readers: dict[Key, list[Key]] = {}
        for key, reads in cone.items():
            for dep in reads:
                readers.setdefault(dep, []).append(key)
                if self.dirty.get(dep) or dep in self.changes:
                    seeds.add(key)
        reached: set[Key] = set()
        work = list(seeds)
        while work:
            key = work.pop()
            if key in reached:
                continue
            reached.add(key)
            work.extend(reader for reader in readers.get(key, ()) if reader in cone)
        for key in cone:
            self.dirty[key] = key in reached

    # -- values ---------------------------------------------------------------------------------
    def value(self, key: Key) -> Any:
        """The value of one cell, computing whatever it needs first -- iteratively, not by
        recursion: a running balance down 50,000 rows is 50,000 deep."""
        if key in self.memo:
            return self.memo[key]
        stack: list[Key] = [key]
        visiting: set[Key] = set()
        while stack:
            current = stack[-1]
            if current in self.memo:
                stack.pop()
                continue
            immediate = self._immediate(current)
            if immediate is not None:
                self._settle(current, *immediate)
                stack.pop()
                continue
            plan = self.plans[current]
            pending = [
                dep
                for dep in self._read_cells(plan.inputs, include_changes=False)
                if dep not in self.memo and self._immediate(dep) is None
            ]
            if not pending:
                visiting.discard(current)
                stack.pop()
                self._compute(current, plan)
                continue
            visiting.add(current)
            cyclic = next((dep for dep in pending if dep in visiting), None)
            if cyclic is not None:
                reason = f"circular reference through {_qualified(*cyclic)}"
                self._settle(current, _Unknown(reason), "unknown")
                visiting.discard(current)
                stack.pop()
                continue
            stack.extend(pending)
        return self.memo[key]

    def _immediate(self, key: Key) -> tuple[Any, Status] | None:
        """A value that needs no computing -- a change, an input, a saved result no change
        reaches, or a known unknown -- or ``None`` when the cell must be computed."""
        if key in self.changes:
            return self.changes[key], "changed"
        cell = self._cell(key)
        if cell is None:
            return self.library.EMPTY, "input"
        if not cell.computed:
            return _excel_value(cell.value, self.library), "input"
        if not self.is_dirty(key):
            return _excel_value(cell.cached_value, self.library), "saved"
        if self.stopped is not None:
            return _Unknown(f"not computed: {self.stopped}"), "unknown"
        if cell.formula is None and cell.array_master is None:
            return _Unknown("its formula text is not stored in this file format"), "unknown"
        plan = self._plan(key)
        if plan.blocked is not None:
            return _Unknown(plan.blocked), "unknown"
        return None

    def _settle(self, key: Key, value: Any, status: Status) -> None:
        self.memo[key] = value
        self.status[key] = status

    def _out_of_time(self) -> bool:
        if self.stopped is None and time.monotonic() > self.deadline:
            self.stopped = "the time limit was reached"
        return self.stopped is not None

    def _compute(self, key: Key, plan: _Plan) -> None:
        if self.stopped is None and self.evaluations >= self.max_evaluations:
            self.stopped = f"the {self.max_evaluations:,}-evaluation limit was reached"
        if self._out_of_time():
            self._settle(key, _Unknown(f"not computed: {self.stopped}"), "unknown")
            return
        cell = self._cell(key)
        assert cell is not None
        if cell.array_master is not None:
            self._settle(key, *self._spilled(key, cell))
            return
        arguments: list[Any] = []
        for region in plan.inputs:
            values = self._range_values(region)
            if isinstance(values, _Unknown):
                self._settle(key, values, "unknown")
                return
            arguments.append(values)
        self.evaluations += 1
        value = self._run(plan, arguments)
        if isinstance(value, _Unknown):
            self._settle(key, value, "unknown")
            return
        if plan.volatile:
            self.volatile.add(_qualified(*key))
        if cell.array_range and cell.array_range != cell.coordinate:
            self.arrays[key] = value
        self._settle(key, _first(value), "recalculated")

    def _run(self, plan: _Plan, arguments: list[Any]) -> Any:
        if plan.fast is not None:
            try:
                return plan.fast(*arguments)
            except ExcelError as error:
                return self.library.Error(error.code)
            except Fallback:
                pass
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                value = plan.function(*arguments)
        except Exception as exc:  # an unimplemented function raises inside the library
            return _Unknown(f"the evaluator cannot compute it ({_brief(exc)})")
        if isinstance(_first(value), complex):
            # A negative number to a fractional power: Excel's #NUM!, not a complex number.
            return self.library.Error("#NUM!")
        return value

    def _spilled(self, key: Key, cell: CellValue) -> tuple[Any, Status]:
        """A cell inside an array formula's extent: its element of the master's new result."""
        master = (key[0], cell.array_master or "")
        master_value = self.memo.get(master)
        if isinstance(master_value, _Unknown):
            return master_value, "unknown"
        result = self.arrays.get(master)
        if result is None:
            return _Unknown("its array formula was not recomputed"), "unknown"
        origin = A1Range.parse(key[0], master[1])
        rows, columns = cell.row - origin.min_row, cell.column - origin.min_col
        grid = _grid(result)
        if rows < len(grid) and columns < len(grid[rows]):
            return grid[rows][columns], "recalculated"
        return self.library.EMPTY, "recalculated"

    def _range_values(self, region: A1Range) -> Any:
        """What the evaluator is fed for one input: a value, or a grid for a range."""
        if region.cell_count == 1:
            value = self.value((region.sheet_name, region.a1))
            if isinstance(value, _Unknown):
                return _unknown_input(region.sheet_name, region.a1, value)
            return value
        import numpy as np

        grid = np.empty(
            (region.max_row - region.min_row + 1, region.max_col - region.min_col + 1),
            dtype=object,
        )
        grid.fill(self.library.EMPTY)
        sheet = self.sheets.get(region.sheet_name)
        if sheet is not None:
            for coordinate in _coordinates(sheet, region, computed_only=False):
                value = self.value((sheet.name, coordinate))
                if isinstance(value, _Unknown):
                    return _unknown_input(sheet.name, coordinate, value)
                cell = sheet.cells[coordinate]
                grid[cell.row - region.min_row, cell.column - region.min_col] = value
        for (sheet_name, coordinate), value in self.changes.items():
            spot = A1Range.parse(sheet_name, coordinate)
            if region.intersects(spot):
                grid[spot.min_row - region.min_row, spot.min_col - region.min_col] = value
        return grid

    # -- plans ----------------------------------------------------------------------------------
    def _plan(self, key: Key) -> _Plan:
        if key in self.plans:
            return self.plans[key]
        cell = self._cell(key)
        assert cell is not None
        if cell.array_master is not None:
            # A cell in an array formula's extent waits on its master, then takes its element.
            plan = _Plan(function=None, inputs=(A1Range.parse(key[0], cell.array_master),))
        elif cell.formula is None:
            plan = _Plan(function=None, inputs=(), blocked="its formula text is not stored")
        else:
            plan = self._compile(key[0], cell)
        self.plans[key] = plan
        return plan

    def _compile(self, sheet: str, cell: CellValue) -> _Plan:
        context = self.model.formula_contexts.get(sheet)
        if context is None:
            return _Plan(None, (), blocked="the workbook was read without formula resolution")
        formula = cell.formula or ""
        rewritten = _rewrite(formula, context, cell.row, self.operands)
        if isinstance(rewritten, _Unknown):
            # Its static edges still say whether a change reaches it.
            static = parse_formula(formula, context, row=cell.row)
            unbounded = any(gap.reason in _UNBOUNDED for gap in static.unresolved)
            edges = tuple(_edge_ranges(static.references))
            return _Plan(None, edges, blocked=rewritten.reason, unbounded=unbounded)
        function = self.compiled.get(rewritten.template)
        if function is None:
            try:
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    function = self.library.Parser().ast(rewritten.template)[1].compile()
            except Exception as exc:  # the library raises its own and builtin errors alike
                reason = f"the evaluator cannot read this formula ({_brief(exc)})"
                return _Plan(None, tuple(rewritten.references.values()), blocked=reason)
            self.compiled[rewritten.template] = function
        inputs: list[A1Range] = []
        for name in function.inputs:
            region = rewritten.references.get(str(name).upper())
            if region is None:
                reason = f"it reads {name!r}, which is not a range of this workbook"
                return _Plan(None, tuple(rewritten.references.values()), blocked=reason)
            if region.cell_count > self.max_range_cells:
                where = _qualified(region.sheet_name, region.a1)
                reason = (
                    f"it reads {region.cell_count:,} cells in {where}, over the "
                    f"{self.max_range_cells:,}-cell limit"
                )
                return _Plan(None, tuple(rewritten.references.values()), blocked=reason)
            inputs.append(region)
        fast = None
        if all(region.cell_count == 1 for region in inputs):
            if rewritten.template not in self.fast:
                names = [str(name) for name in function.inputs]
                self.fast[rewritten.template] = compile_arithmetic(rewritten.template, names)
            fast = self.fast[rewritten.template]
        return _Plan(function, tuple(inputs), volatile=rewritten.volatile, fast=fast)

    def _read_cells(self, inputs: Iterable[A1Range], *, include_changes: bool) -> list[Key]:
        """The computed cells inside the rectangles (and, if asked, the changed ones): what a
        formula's value can depend on beyond plain inputs."""
        found: list[Key] = []
        for region in inputs:
            sheet = self.sheets.get(region.sheet_name)
            if sheet is not None:
                for coordinate in _coordinates(sheet, region, computed_only=True):
                    key = (sheet.name, coordinate)
                    if key not in self.changes:
                        found.append(key)
            if include_changes:
                found.extend(
                    key
                    for key in self.changes
                    if key[0] == region.sheet_name
                    and region.intersects(A1Range.parse(key[0], key[1]))
                )
        return found

    def _cell(self, key: Key) -> CellValue | None:
        sheet = self.sheets.get(key[0])
        return None if sheet is None else sheet.cells.get(key[1])


def _unknown_input(sheet: str, coordinate: str, value: _Unknown) -> _Unknown:
    return _Unknown(f"it reads {_qualified(sheet, coordinate)}, which is unknown ({value.reason})")


def _edge_ranges(references: Iterable[Any]) -> Iterable[A1Range]:
    for edge in references:
        try:
            yield A1Range.parse(edge.reference.sheet_name, edge.reference.a1_range)
        except ValueError:
            continue


#: Why a formula may read cells no static edge names: any change may reach it.
_UNBOUNDED = frozenset(
    {
        UnresolvedReason.INDIRECT,
        UnresolvedReason.VOLATILE_OFFSET,
        UnresolvedReason.UNSUPPORTED_FUNCTION,
        UnresolvedReason.DYNAMIC_ARRAY,
    }
)


# -------------------------------------------------------------------------------------------------
# Rewriting a formula into plain A1 references the evaluator can feed
# -------------------------------------------------------------------------------------------------
_ANCHORARRAY_RE = re.compile(
    r"(?i)(?:_xlfn\.)?ANCHORARRAY\(\s*((?:(?:'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?"
    r"\$?[A-Za-z]{1,3}\$?\d{1,7})\s*\)"
)


@dataclass(frozen=True, slots=True)
class _Rewrite:
    """A formula as the evaluator compiles it: each distinct reference replaced by a placeholder
    name (``_REF0``...), so every row of a filled-down column shares one compiled function, and
    the rectangle each placeholder stands for in this cell."""

    template: str
    references: dict[str, A1Range]
    volatile: bool


#: A plain A1 cell or rectangle, optionally sheet-qualified: resolved without the full parser.
_PLAIN_REF_RE = re.compile(
    r"(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<bare>[A-Za-z0-9_.]+))!)?"
    r"(?P<ref>\$?[A-Za-z]{1,3}\$?\d{1,7}(?::\$?[A-Za-z]{1,3}\$?\d{1,7})?)"
)

Operands = dict[tuple[str, str, int], "A1Range | _Unknown"]


def _rewrite(
    formula: str, context: FormulaContext, row: int, cache: Operands
) -> _Rewrite | _Unknown:
    """``formula`` with every reference replaced by a placeholder for the A1 rectangle it reads,
    and whether it calls a volatile function; or why it cannot be evaluated. ``cache`` remembers
    resolved operands across the formulas of one calculation."""
    text = formula[1:] if formula.startswith("=") else formula
    spilled = _replace_spills(text, context, row)
    if isinstance(spilled, _Unknown):
        return spilled
    try:
        tokens = Tokenizer("=" + spilled).items
    except TokenizerError as exc:
        return _Unknown(f"the formula cannot be read ({exc})")
    bound = _bound_names(tokens)
    volatile = False
    parts: list[str] = []
    placeholders: dict[A1Range, str] = {}
    for token in tokens:
        value = token.value
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            name = value[:-1].strip()
            if name.upper().rsplit(".", 1)[-1] in VOLATILE_FUNCTIONS:
                volatile = True
            if ":" in name:
                return _Unknown("a range bounded by a function result is decided at calculation")
            gaps = _function_gaps(name)
            if gaps:
                return _Unknown(f"it uses {gaps[0].reference_text}: {gaps[0].detail}")
        elif token.type == Token.OPERAND and token.subtype == Token.RANGE:
            stripped = value.strip()
            if stripped.lower() in bound or stripped.lower().startswith("_xlpm."):
                # A LET/LAMBDA variable; on disk it carries `_xlpm.`, which the evaluator lacks.
                parts.append(re.sub(r"(?i)_xlpm\.", "", value))
                continue
            region = _operand(stripped, context, row, cache)
            if isinstance(region, _Unknown):
                return region
            value = placeholders.setdefault(region, f"_REF{len(placeholders)}")
        parts.append(value)
    return _Rewrite(
        template="=" + "".join(parts),
        references={name: region for region, name in placeholders.items()},
        volatile=volatile,
    )


def _operand(text: str, context: FormulaContext, row: int, cache: Operands) -> A1Range | _Unknown:
    """The rectangle one reference operand reads, by this package's resolver (names, tables,
    whole columns), with a fast path for a plain A1 reference."""
    this_row = "@" in text or "this row" in text.lower()
    key = (context.sheet_name, text, row if this_row else 0)
    if key in cache:
        return cache[key]
    resolved: A1Range | _Unknown
    plain = _PLAIN_REF_RE.fullmatch(text)
    sheet = None
    if plain is not None:
        quoted, bare = plain.group("quoted"), plain.group("bare")
        sheet = quoted.replace("''", "'") if quoted is not None else bare or context.sheet_name
    if plain is not None and sheet in context.known_sheets:
        try:
            resolved = A1Range.parse(sheet, plain.group("ref").replace("$", "").upper())
        except ValueError:
            resolved = _Unknown(f"it reads {text}, which is not an A1 reference")
    else:
        parsed = parse_formula(f"={text}", context, row=row)
        blocking = [gap for gap in parsed.unresolved if gap.reason in _BLOCKING_REASONS]
        if blocking:
            resolved = _Unknown(f"it reads {blocking[0].reference_text}: {blocking[0].detail}")
        elif len(parsed.references) != 1:
            resolved = _Unknown(f"it reads {text}, which is not one rectangle")
        else:
            reference = parsed.references[0].reference
            try:
                resolved = A1Range.parse(reference.sheet_name, reference.a1_range)
            except ValueError:
                resolved = _Unknown(f"it reads {text}, which is not a rectangle")
    cache[key] = resolved
    return resolved


def _replace_spills(text: str, context: FormulaContext, row: int) -> str | _Unknown:
    """Replace each spill reference (``B2#``, ``_xlfn.ANCHORARRAY(B2)``) with the rectangle its
    array formula last covered, which is what the evaluator can read."""
    outside = _blank(text, [match.span() for match in _STRING_RE.finditer(text)])
    for match in reversed(list(_ANCHORARRAY_RE.finditer(outside))):
        text = f"{text[: match.start()]}{text[match.start(1) : match.end(1)]}#{text[match.end() :]}"
    outside = _blank(text, [match.span() for match in _STRING_RE.finditer(text)])
    for match in reversed(list(_SPILL_RE.finditer(outside))):
        anchor = _SPILL_ANCHOR_RE.search(outside[: match.start()])
        if anchor is None:
            return _Unknown("it reads a spill range whose anchor is not a cell")
        parsed = parse_formula(f"={text[anchor.start() : match.start()]}#", context, row=row)
        spills = [edge.reference for edge in parsed.references if edge.reference.kind == "spill"]
        if len(spills) != 1:
            return _Unknown("it reads a spill range with no saved extent")
        extent = _qualified(spills[0].sheet_name, spills[0].a1_range, always_quote=True)
        text = f"{text[: anchor.start()]}{extent}{text[match.end() :]}"
        outside = _blank(text, [found.span() for found in _STRING_RE.finditer(text)])
    return text


# -------------------------------------------------------------------------------------------------
# Which formulas a change might reach (a cheap superset)
# -------------------------------------------------------------------------------------------------
def _reachable_entries(model: WorkbookModel, changes: Mapping[Key, Any]) -> set[str]:
    """The formula entries (by node id) a change might reach, through the static edges.

    A superset, cheaply: an entry of repeated formulas carries edges widened over all its rows,
    and a formula that may read cells no edge names (``INDIRECT``...) counts as reached. The engine
    decides exactly, per cell, only within this set.
    """
    if not changes:
        return set()
    entries = [entry for sheet in model.sheets for entry in sheet.formulas]
    edges: dict[str, list[tuple[FormulaEntry, Any]]] = {}
    for entry in entries:
        for reference in entry.references:
            if reference.row_span is not None and reference.column_span is not None:
                edges.setdefault(reference.sheet_name, []).append((entry, reference))
    reached: set[str] = set()
    frontier = [A1Range.parse(sheet, coordinate) for sheet, coordinate in changes]
    for entry in entries:
        if any(gap.reason in _UNBOUNDED for gap in entry.unresolved_references):
            reached.add(entry.node_id)
            frontier.append(_entry_extent(entry))
    while frontier:
        span = frontier.pop()
        for entry, reference in edges.get(span.sheet_name, ()):
            if entry.node_id in reached:
                continue
            rows, columns = reference.row_span, reference.column_span
            if (
                rows["gte"] <= span.max_row
                and span.min_row <= rows["lte"]
                and columns["gte"] <= span.max_col
                and span.min_col <= columns["lte"]
            ):
                reached.add(entry.node_id)
                frontier.append(_entry_extent(entry))
    return reached


def _entry_extent(entry: FormulaEntry) -> A1Range:
    if entry.array_range:
        try:
            return A1Range.parse(entry.sheet_name, entry.array_range)
        except ValueError:
            pass
    return entry.a1_range


# -------------------------------------------------------------------------------------------------
# Inputs and values
# -------------------------------------------------------------------------------------------------
def calculation_available() -> bool:
    """Whether the ``calc`` extra is installed, without importing it."""
    import importlib.util

    return importlib.util.find_spec("formulas") is not None


def _load_formulas() -> Any:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            import formulas  # type: ignore[import-untyped]
            import schedula  # type: ignore[import-untyped]
            from formulas.tokens.operand import XlError  # type: ignore[import-untyped]
    except ImportError as exc:
        raise CalcUnavailable(
            "formula evaluation needs the calc extra: pip install 'excel-rag[calc]' "
            "(uv sync --extra calc)"
        ) from exc

    class _Library:
        Parser = formulas.Parser
        EMPTY = schedula.EMPTY
        Error = XlError

    return _Library


def _parse_changes(model: WorkbookModel, changes: Mapping[str, Any]) -> dict[tuple[str, str], Any]:
    library = _load_formulas()
    parsed: dict[tuple[str, str], Any] = {}
    for target, raw in changes.items():
        region = _parse_target(model, target)
        if region.cell_count != 1:
            raise CalcInputError(f"a change sets one cell, not {target!r}")
        if isinstance(raw, str) and raw.strip().startswith("="):
            raise CalcInputError(f"a change sets a value, not a formula: {target}={raw}")
        parsed[(region.sheet_name, region.a1)] = _excel_value(_typed(raw), library)
    return parsed


def _typed(raw: Any) -> Any:
    """A typed-in string as Excel would take it: a number, a boolean, an error or text."""
    if not isinstance(raw, str):
        return raw
    text = raw.strip()
    if text.upper() in ("TRUE", "FALSE"):
        return text.upper() == "TRUE"
    if text.upper() in ERROR_VALUES:
        return text.upper()
    cleaned = text.replace(",", "")
    percent = cleaned.endswith("%")
    try:
        number = float(cleaned[:-1] if percent else cleaned)
    except ValueError:
        return raw
    if not math.isfinite(number):
        return raw
    number = number / 100 if percent else number
    return int(number) if number.is_integer() and not percent and "." not in cleaned else number


def _excel_value(value: Any, library: Any) -> Any:
    """A reader or caller value as the evaluator takes it."""
    if value is None or value == "":
        return library.EMPTY
    if isinstance(value, str) and value in ERROR_VALUES:
        return library.Error(value)
    if isinstance(value, datetime | date | dtime):
        from openpyxl.utils.datetime import to_excel  # type: ignore[import-untyped]

        return to_excel(value)
    return value


def _plain(value: Any) -> Any:
    """An evaluator value as plain Python: numbers, text, booleans, error strings, ``None``."""
    if isinstance(value, _Unknown):
        return None
    if value is None:
        return None
    type_name = type(value).__name__
    if type_name == "XlError":
        return str(value)
    if type_name == "Token" and str(value) == "empty":
        return None
    if hasattr(value, "shape") and getattr(value, "shape", None) == ():
        return _plain(value.item() if hasattr(value, "item") else value[()])
    if hasattr(value, "tolist") and getattr(value, "ndim", 0) >= 1:
        return [[_plain(item) for item in row] for row in _grid(value)]
    if hasattr(value, "item") and type(value).__module__ == "numpy":
        return value.item()
    if isinstance(value, float) and value.is_integer() and abs(value) < 2**53:
        return int(value)
    return value


def _first(value: Any) -> Any:
    """The top-left value of an array result, which is what its cell shows."""
    if hasattr(value, "ndim") and value.ndim >= 1:
        grid = _grid(value)
        return grid[0][0] if grid and grid[0] else None
    if hasattr(value, "shape") and value.shape == ():
        return value[()]
    return value


def _grid(value: Any) -> list[list[Any]]:
    if hasattr(value, "ndim"):
        if value.ndim == 0:
            return [[value[()]]]
        if value.ndim == 1:
            return [list(value)]
        return [list(row) for row in value]
    return [[value]]


def _differs(result: CellResult) -> bool:
    if result.status != "recalculated" or result.saved is None:
        return False
    if isinstance(result.value, int | float) and isinstance(result.saved, int | float):
        if isinstance(result.value, bool) or isinstance(result.saved, bool):
            return result.value != result.saved
        return not math.isclose(result.value, result.saved, rel_tol=1e-9, abs_tol=1e-9)
    return bool(result.value != result.saved)


# -------------------------------------------------------------------------------------------------
# Targets
# -------------------------------------------------------------------------------------------------
def _as_list(targets: str | Sequence[str]) -> list[str]:
    return [targets] if isinstance(targets, str) else list(targets)


def _parse_target(model: WorkbookModel, text: str) -> A1Range:
    match = re.fullmatch(
        r"\s*(?:'(?P<quoted>(?:[^']|'')+)'|(?P<bare>[^'!]+))!(?P<ref>[^!]+?)\s*", text
    )
    if match is None:
        raise CalcInputError(f"name a cell or range with its sheet, like Sheet1!B4: {text!r}")
    sheet_text = (match.group("quoted") or "").replace("''", "'") or (match.group("bare") or "")
    sheet = next(
        (s.name for s in model.sheets if s.name.casefold() == sheet_text.strip().casefold()), None
    )
    if sheet is None:
        names = ", ".join(repr(s.name) for s in model.sheets)
        raise CalcInputError(f"no sheet named {sheet_text!r}; the sheets are {names}")
    try:
        return A1Range.parse(sheet, match.group("ref").replace("$", "").upper())
    except ValueError as exc:
        raise CalcInputError(f"not an A1 cell or range: {text!r}") from exc


def _cells(model: WorkbookModel, rectangles: Iterable[A1Range]) -> list[tuple[str, str]]:
    """The populated (or changed) cells of the targets, in reading order, once each."""
    sheets = {sheet.name: sheet for sheet in model.sheets}
    found: dict[tuple[str, str], None] = {}
    for region in rectangles:
        sheet = sheets[region.sheet_name]
        if region.cell_count == 1:
            found[(sheet.name, region.a1)] = None
            continue
        for coordinate in _coordinates(sheet, region, computed_only=False):
            found[(sheet.name, coordinate)] = None
    return list(found)


def _coordinates(sheet: SheetModel, region: A1Range, *, computed_only: bool) -> list[str]:
    """The populated cells of ``sheet`` inside ``region``, in reading order."""
    if region.cell_count <= 4 * max(1, len(sheet.cells)):
        coordinates = (
            f"{column_letter(column)}{row}"
            for row in range(region.min_row, region.max_row + 1)
            for column in range(region.min_col, region.max_col + 1)
        )
        cells = (sheet.cells.get(coordinate) for coordinate in coordinates)
        chosen = [cell for cell in cells if cell is not None]
    else:
        chosen = sorted(
            (
                cell
                for cell in sheet.cells.values()
                if region.min_row <= cell.row <= region.max_row
                and region.min_col <= cell.column <= region.max_col
            ),
            key=_order,
        )
    return [cell.coordinate for cell in chosen if cell.computed or not computed_only]


def _order(cell: CellValue) -> tuple[int, int]:
    return (cell.row, cell.column)


def _qualified(sheet: str, a1: str, *, always_quote: bool = False) -> str:
    plain = sheet.replace("_", "").isalnum() and not always_quote
    return f"{sheet}!{a1}" if plain else f"'{sheet.replace(chr(39), chr(39) * 2)}'!{a1}"


def _brief(exc: Exception) -> str:
    text = " ".join(str(exc).split())
    return text[:160] if text else type(exc).__name__


# -------------------------------------------------------------------------------------------------
# Text for a person or a model
# -------------------------------------------------------------------------------------------------
def format_calculation(calculation: Calculation, *, limit: int = 200) -> str:
    """The calculation as plain text: the changes, each target cell's value and where it came
    from, and anything unknown or cut short."""
    lines: list[str] = []
    if calculation.changes:
        changed = ", ".join(
            f"{_qualified(sheet, coordinate)} = {_show(value)}"
            for sheet, coordinate, value in calculation.changes
        )
        lines.append(f"With {changed}:")
    shown = calculation.results[:limit]
    for result in shown:
        lines.append(f"- {_result_line(result)}")
    if len(calculation.results) > limit:
        lines.append(f"({len(calculation.results) - limit} more cell(s) not shown.)")
    if not calculation.results:
        lines.append("No populated cells in that range.")
    if calculation.mismatches:
        lines.append(
            f"{len(calculation.mismatches)} recomputed value(s) differ from Excel's saved ones."
        )
    if calculation.volatile:
        lines.append(
            "Volatile (NOW, TODAY, RAND...): "
            + ", ".join(calculation.volatile[:20])
            + " -- values are as of this calculation."
        )
    if calculation.stopped:
        lines.append(f"Stopped early: {calculation.stopped}.")
    lines.append(
        f"{calculation.recalculated} formula evaluation(s) by excel-rag; 'saved' values are "
        "Excel's own, which no change reaches."
    )
    return "\n".join(lines)


def format_check(calculation: Calculation, *, limit: int = 50) -> str:
    """A full recalculation compared with Excel's saved values: how many agree, which differ, and
    why the rest could not be computed."""
    compared = [
        result
        for result in calculation.results
        if result.status == "recalculated"
        and result.saved is not None
        and result.qualified not in calculation.volatile
    ]
    unknown = [result for result in calculation.results if result.status == "unknown"]
    agreeing = len(compared) - len(calculation.mismatches)
    lines = [
        f"Recomputed {len(calculation.results)} formula cell(s): {agreeing} agree with Excel's "
        f"saved values, {len(calculation.mismatches)} differ, {len(unknown)} could not be computed."
    ]
    if calculation.mismatches:
        lines.append("Differ:")
        lines.extend(f"- {_result_line(result)}" for result in calculation.mismatches[:limit])
    if unknown:
        reasons: dict[str, int] = {}
        for result in unknown:
            reason = (result.reason or "").split(" (", 1)[0]
            reasons[reason] = reasons.get(reason, 0) + 1
        lines.append("Could not be computed:")
        lines.extend(
            f"- {count} cell(s): {reason}"
            for reason, count in sorted(reasons.items(), key=lambda item: -item[1])[:limit]
        )
    if calculation.volatile:
        lines.append(f"{len(calculation.volatile)} volatile cell(s) (NOW, RAND...) not compared.")
    if calculation.stopped:
        lines.append(f"Stopped early: {calculation.stopped}.")
    return "\n".join(lines)


def _result_line(result: CellResult) -> str:
    where = result.qualified
    if result.status == "unknown":
        saved = f" (Excel last saved {_show(result.saved)})" if result.saved is not None else ""
        return f"{where}: unknown{saved} -- {result.reason}"
    if result.status == "recalculated":
        if result.saved is not None and _differs(result):
            return (
                f"{where}: {_show(result.value)} (recalculated; Excel saved {_show(result.saved)})"
            )
        return f"{where}: {_show(result.value)} (recalculated)"
    return f"{where}: {_show(result.value)} ({result.status})"


def _show(value: Any) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return f"{value:.10g}"
    if isinstance(value, list):
        return "[" + "; ".join(", ".join(_show(item) for item in row) for row in value) + "]"
    if isinstance(value, str) and value not in ERROR_VALUES:
        return f'"{value}"'
    return str(value)


__all__ = [
    "MAX_EVALUATIONS",
    "MAX_RANGE_CELLS",
    "TIMEOUT_SECONDS",
    "CalcInputError",
    "CalcUnavailable",
    "Calculation",
    "CellResult",
    "calculate",
    "calculation_available",
    "check_saved_values",
    "format_calculation",
    "format_check",
]
