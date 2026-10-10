"""A small planning workbook and the questions it answers, built in code -- no binary committed.

It is shaped like the workbooks the service is for: an assumptions sheet with named inputs, a
monthly actuals table, a forecast whose formulas read both, a notes block that states the currency,
and a headcount table. Each question names the sheet and rectangle a correct answer cites; several
are worded with no word in common with the cell text, and some are not in English, because those
are the questions lexical search alone cannot answer.
"""

from __future__ import annotations

from pathlib import Path

from .cases import EvalCase, Expected

SAMPLE_WORKBOOK_ID = "planning"

_REGIONS = ("North", "South", "East", "West")
_MONTHS = ("2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06")


def build_sample_workbook(directory: Path) -> Path:
    """Write ``planning.xlsx`` into ``directory`` and return its path."""
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.workbook.defined_name import DefinedName  # type: ignore[import-untyped]

    wb = openpyxl.Workbook()
    assumptions = wb.active
    assumptions.title = "Assumptions"
    assumptions["A1"] = "Planning assumptions"
    assumptions.append([])
    assumptions.append(["Parameter", "Value"])
    for label, value in (
        ("Growth rate", 0.05),
        ("Tax rate", 0.21),
        ("EUR/USD exchange rate", 1.08),
        ("Discount rate", 0.08),
    ):
        assumptions.append([label, value])
    for row in (4, 5, 7):
        assumptions.cell(row=row, column=2).number_format = "0%"
    wb.defined_names["GrowthRate"] = DefinedName("GrowthRate", attr_text="Assumptions!$B$4")
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Assumptions!$B$5")

    actuals = wb.create_sheet("Actuals")
    actuals.append(["Month", "Region", "Revenue", "Cost", "Units sold"])
    for month_index, month in enumerate(_MONTHS):
        for region_index, region in enumerate(_REGIONS):
            base = 1000 + 150 * region_index + 40 * month_index
            actuals.append([month, region, base, round(base * 0.62), 20 + region_index])

    forecast = wb.create_sheet("Forecast")
    forecast["A1"] = "Revenue forecast 2027"
    forecast.append([])
    forecast.append(["Quarter", "Revenue", "Cost", "EBITDA", "Net income"])
    for index, quarter in enumerate(("Q1", "Q2", "Q3", "Q4"), start=4):
        forecast.append(
            [
                quarter,
                f"=SUM(Actuals!C2:C25)/2*(1+GrowthRate)^{index - 3}",
                f"=SUM(Actuals!D2:D25)/2*(1+GrowthRate)^{index - 3}",
                f"=B{index}-C{index}",
                f"=D{index}*(1-TaxRate)",
            ]
        )
    forecast["A10"] = "Notes:"
    forecast["A11"] = "All amounts are reported in US dollars at constant currency."
    forecast["A12"] = "Quarterly revenue grows by the planning growth rate."

    headcount = wb.create_sheet("Headcount")
    headcount.append(["Department", "Employees", "Average salary"])
    for department, employees, salary in (
        ("Engineering", 42, 98000),
        ("Sales", 25, 72000),
        ("Finance", 8, 81000),
        ("Operations", 15, 64000),
    ):
        headcount.append([department, employees, salary])

    path = directory / "planning.xlsx"
    wb.save(path)
    return path


def _case(identifier: str, question: str, *expected: tuple[str, str]) -> EvalCase:
    return EvalCase(
        id=identifier,
        question=question,
        workbook_id=SAMPLE_WORKBOOK_ID,
        expected=tuple(Expected(sheet, a1) for sheet, a1 in expected),
    )


#: Questions over :func:`build_sample_workbook`, each with the rectangle a correct answer cites.
SAMPLE_CASES: tuple[EvalCase, ...] = (
    _case("growth-rate", "What growth rate does the plan assume?", ("Assumptions", "A3:B7")),
    _case("tax", "Which tax rate is applied to profit?", ("Assumptions", "A3:B7")),
    _case("fx", "What exchange rate converts euros to dollars?", ("Assumptions", "A3:B7")),
    _case("revenue-by-region", "Revenue by region per month", ("Actuals", "A1:E25")),
    _case("units", "How many units were sold?", ("Actuals", "E1:E25")),
    _case(
        "forecast-revenue",
        "How is projected revenue calculated?",
        ("Forecast", "A3:E7"),
    ),
    _case("ebitda", "What is the EBITDA for each quarter?", ("Forecast", "D3:D7")),
    _case("net-income", "net income after tax", ("Forecast", "E3:E7")),
    _case("currency", "What currency are the amounts in?", ("Forecast", "A10:A12")),
    _case("staff", "How many people work in engineering?", ("Headcount", "A1:C5")),
    _case("pay", "average pay by department", ("Headcount", "A1:C5")),
    # No word in common with the cells: only a semantic retriever can find these.
    _case("payroll", "What does each team cost in wages?", ("Headcount", "A1:C5")),
    _case("profit-outlook", "What profit do we expect next year?", ("Forecast", "A3:E7")),
    # Not in English.
    _case("es-currency", "¿En qué moneda están los importes?", ("Forecast", "A10:A12")),
    _case("fr-headcount", "Combien d'employés par département ?", ("Headcount", "A1:C5")),
    _case("de-growth", "Welche Wachstumsrate wird angenommen?", ("Assumptions", "A3:B7")),
)


__all__ = ["SAMPLE_CASES", "SAMPLE_WORKBOOK_ID", "build_sample_workbook"]
