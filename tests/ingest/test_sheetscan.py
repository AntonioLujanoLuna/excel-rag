"""The streaming scan must give every formula cell the value openpyxl's data_only load gives it."""

from __future__ import annotations

import inspect
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from fixtures import make_fixtures as mk

from excel_rag.workbook import reader
from excel_rag.workbook.errors import WorkbookError
from excel_rag.workbook.reader import read_workbook
from excel_rag.workbook.sheetscan import scan_sheet

# openpyxl reads neither CSV nor .xlsb, so they have no data_only load to compare against.
_SKIP = {"corrupt_file", "copy_as", "large_region", "csv_file", "xlsb_workbook"}
_BUILDERS = sorted(
    name
    for name, member in inspect.getmembers(mk, inspect.isfunction)
    if not name.startswith("_") and name not in _SKIP and member.__module__ == mk.__name__
)


def _saved_values_workbook(directory: Path) -> Path:
    """Formula cells with each saved result type: number, date, duration, text, bool, error."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Types"
    formulas = {
        "A1": ("=1+1", None),
        "A2": ("=TODAY()", "yyyy-mm-dd"),
        "A3": ("=1/24", "[h]:mm:ss"),
        "A4": ('="a"&"b"', None),
        "A5": ("=TRUE()", None),
        "A6": ("=1/0", None),
        "A7": ("=2.5*2", None),
        "A8": ("=A1", None),
    }
    for coordinate, (formula, number_format) in formulas.items():
        ws[coordinate] = formula
        if number_format:
            ws[coordinate].number_format = number_format
    base = directory / "types_base.xlsx"
    wb.save(base)
    saved = {
        "A1": ("n", "2"),
        "A2": ("n", "46000"),
        "A3": ("n", "4.1666666666666664E-2"),
        "A4": ("str", "ab"),
        "A5": ("b", "1"),
        "A6": ("e", "#DIV/0!"),
        "A7": ("n", "5"),
    }

    def edit(data: bytes) -> bytes:
        text = data.decode()
        for coordinate, (kind, value) in saved.items():
            text = text.replace(
                f'<c r="{coordinate}"',
                f'<c r="{coordinate}" t="{kind}"' if kind != "n" else f'<c r="{coordinate}"',
                1,
            )
            head, sep, tail = text.partition(f'<c r="{coordinate}"')
            tail = tail.replace("<v />", f"<v>{value}</v>", 1)
            text = head + sep + tail
        return text.encode()

    return mk._rewrite_zip(base, directory / "types.xlsx", edits={"xl/worksheets/sheet1.xml": edit})


def _openpyxl_saved(path: Path) -> dict[tuple[str, str], object]:
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    return {
        (ws.title, cell.coordinate): values[ws.title][cell.coordinate].value
        for ws in formulas.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }


def _reader_saved(path: Path) -> dict[tuple[str, str], object]:
    raw = read_workbook(path)
    return {
        (sheet.name, cell.coordinate): cell.cached_value
        for sheet in raw.sheets
        for cell in sheet.cells.values()
        if cell.formula is not None
    }


@pytest.mark.parametrize("builder", [*_BUILDERS, "saved_types"])
def test_scan_matches_openpyxl_data_only(builder: str, tmp_path: Path) -> None:
    if builder == "saved_types":
        path = _saved_values_workbook(tmp_path)
    else:
        path = getattr(mk, builder)(tmp_path)
    assert _reader_saved(path) == _openpyxl_saved(path)


def test_saved_types_convert_as_openpyxl_does(tmp_path: Path) -> None:
    saved = _reader_saved(_saved_values_workbook(tmp_path))
    assert saved[("Types", "A1")] == 2
    assert isinstance(saved[("Types", "A2")], datetime)
    assert isinstance(saved[("Types", "A3")], timedelta)
    assert saved[("Types", "A4")] == "ab"
    assert saved[("Types", "A5")] is True
    assert saved[("Types", "A6")] == "#DIV/0!"
    assert saved[("Types", "A7")] == 5
    assert saved[("Types", "A8")] is None


def test_the_second_openpyxl_load_is_skipped(tmp_path: Path, monkeypatch) -> None:
    loads: list[bool] = []
    original = reader._load

    def counting(data: bytes, *, data_only: bool):
        loads.append(data_only)
        return original(data, data_only=data_only)

    monkeypatch.setattr(reader, "_load", counting)
    read_workbook(mk.cross_sheet_formula(tmp_path))
    assert loads == [False]


def test_a_shared_string_result_falls_back_to_openpyxl(tmp_path: Path, monkeypatch) -> None:
    path = mk.cached_value(tmp_path)
    loads: list[bool] = []
    original = reader._load

    def counting(data: bytes, *, data_only: bool):
        loads.append(data_only)
        return original(data, data_only=data_only)

    def mark_shared(stream, part):
        scan = scan_sheet(stream, part)
        scan.needs_shared_strings = True
        return scan

    monkeypatch.setattr(reader, "_load", counting)
    monkeypatch.setattr(reader, "scan_sheet", mark_shared)
    assert _reader_saved(path) == _openpyxl_saved(path)
    assert loads == [False, True]


def test_an_entity_declaring_sheet_is_refused(tmp_path: Path) -> None:
    import io

    bomb = (
        b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<sheetData/></worksheet>"
    )
    with pytest.raises(WorkbookError, match="entities"):
        scan_sheet(io.BytesIO(bomb), "sheet1")


def test_dimension_is_read_without_a_tree(tmp_path: Path) -> None:
    import io

    part = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:XFD1048576"/><sheetData><row r="3"><c><f>1+1</f><v>2</v></c>'
        b"</row></sheetData></worksheet>"
    )
    scan = scan_sheet(io.BytesIO(part), "sheet1")
    assert scan.dimension == "A1:XFD1048576"
    assert scan.cached["A3"].text == "2"
