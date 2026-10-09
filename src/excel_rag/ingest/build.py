"""Assemble a :class:`~excel_rag.ingest.canonical.WorkbookModel` from raw sheets.

This is the seam between the reader (bytes -> raw cells), region detection (cells -> regions),
formula extraction (formula text -> edges) and the document builder. Formula cells are grouped by
their normalised pattern first, so a column of 500 identical formulas becomes one *cluster* node and
one summary instead of 500 cell documents -- the design's "group repeated formulas into patterns".
"""

from __future__ import annotations

import re
from collections.abc import Sequence

from openpyxl.formula.tokenizer import (  # type: ignore[import-untyped]
    Token,
    Tokenizer,
    TokenizerError,
)

from ..models import A1Range, Reference, ReferenceKind
from .canonical import (
    CellValue,
    FormulaEntry,
    NamedRange,
    Region,
    SheetModel,
    WorkbookModel,
    cell_node_id,
    column_letter,
    formula_node_id,
    range_node_id,
    sheet_node_id,
)
from .formulas import (
    FormulaContext,
    NamedRangeInfo,
    ParsedFormula,
    ReferenceEdge,
    TableInfo,
    parse_formula,
    resolve_named_range,
)
from .reader import RawSheet, RawWorkbook
from .regions import RegionConfig, detect_regions

#: Row digits of *relative* references are replaced with ``#`` so identical patterns group.
_REF_ROW_RE = re.compile(r"(?<![A-Za-z0-9_])(\$?[A-Za-z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_!])")


def normalize_formula(formula: str) -> str:
    """A pattern key: relative row numbers collapse, absolute and constant parts stay literal.

    Only reference operands are rewritten -- string literals, numbers and function names (``LOG10``)
    are kept verbatim -- so ``="Q1"&A2`` and ``="Q2"&A3`` are different patterns. A formula the
    tokenizer refuses falls back to rewriting the whole text.
    """

    def replace(match: re.Match[str]) -> str:
        column, dollar, _row = match.group(1), match.group(2), match.group(3)
        return f"{column}#" if dollar == "" else match.group(0)

    try:
        tokens = Tokenizer(formula if formula.startswith("=") else f"={formula}").items
    except TokenizerError:
        return _REF_ROW_RE.sub(replace, formula)
    parts = [
        _REF_ROW_RE.sub(replace, token.value)
        if token.type == Token.OPERAND and token.subtype == Token.RANGE
        else token.value
        for token in tokens
    ]
    return ("=" if formula.startswith("=") else "") + "".join(parts)


def _col_range_a1(min_col: int, max_col: int, min_row: int, max_row: int) -> str:
    start = f"{column_letter(min_col)}{min_row}"
    end = f"{column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _expand_cluster_edges(
    edges: Sequence[ReferenceEdge],
    *,
    sheet_name: str,
    representative_row: int,
    group_min_row: int,
    group_max_row: int,
    workbook_id: str,
    version: int,
) -> tuple[Reference, ...]:
    """Broaden a representative formula's relative edges to cover the whole cluster's rows."""
    expanded: list[Reference] = []
    shift = group_min_row - representative_row
    extra = group_max_row - group_min_row
    for edge in edges:
        reference = edge.reference
        if edge.absolute or reference.sheet_name != sheet_name:
            expanded.append(reference)
            continue
        new_min = max(1, edge.min_row + shift)
        new_max = max(new_min, edge.max_row + shift + extra)
        a1 = _col_range_a1(edge.min_col, edge.max_col, new_min, new_max)
        single = new_min == new_max and edge.min_col == edge.max_col
        target = (
            cell_node_id(workbook_id, version, sheet_name, a1)
            if single
            else range_node_id(workbook_id, version, sheet_name, a1)
        )
        expanded.append(
            Reference(
                target_node_id=target,
                sheet_name=sheet_name,
                a1_range=a1,
                kind=ReferenceKind.CELL if single else ReferenceKind.RANGE,
            )
        )
    return tuple(expanded)


def _used_range(sheet: RawSheet) -> A1Range | None:
    if not sheet.cells:
        return None
    rows = [cell.row for cell in sheet.cells.values()]
    cols = [cell.column for cell in sheet.cells.values()]
    a1 = A1Range.parse(
        sheet.name,
        f"{column_letter(min(cols))}{min(rows)}:{column_letter(max(cols))}{max(rows)}",
    )
    return a1


def _build_tables(raw: RawWorkbook) -> dict[str, TableInfo]:
    resolved: dict[str, TableInfo] = {}
    for sheet in raw.sheets:
        for table in sheet.tables:
            if not table.ref:
                continue
            try:
                parsed = A1Range.parse(sheet.name, table.ref)
            except ValueError:
                continue
            info = TableInfo(table.name, sheet.name, parsed, table.columns)
            resolved[table.name] = info
            resolved[table.name.lower()] = info
    return resolved


def _build_named_ranges(raw: RawWorkbook, workbook_id: str, version: int) -> tuple[NamedRange, ...]:
    known = frozenset(sheet.name for sheet in raw.sheets if not sheet.is_macro_sheet)
    sheet_by_index = {index: sheet.name for index, sheet in enumerate(raw.sheets)}
    names = [
        resolve_named_range(
            defined.name,
            defined.attr_text,
            defined.local_sheet_id,
            workbook_id=workbook_id,
            version=version,
            known_sheets=known,
            sheet_by_index=sheet_by_index,
        )
        for defined in raw.defined_names
    ]
    names.sort(key=lambda name: name.name.lower())
    return tuple(names)


def _build_formulas(
    sheet: RawSheet,
    context: FormulaContext,
) -> tuple[FormulaEntry, ...]:
    formula_cells = [cell for cell in sheet.cells.values() if cell.formula is not None]
    groups: dict[str, list[CellValue]] = {}
    for cell in formula_cells:
        groups.setdefault(normalize_formula(cell.formula or ""), []).append(cell)
    entries: list[FormulaEntry] = []
    for pattern in sorted(groups):
        cells = sorted(groups[pattern], key=lambda cell: (cell.row, cell.column))
        parsed: ParsedFormula = parse_formula(cells[0].formula or "", context)
        if len(cells) == 1:
            cell = cells[0]
            entries.append(
                FormulaEntry(
                    node_id=formula_node_id(
                        context.workbook_id, context.version, sheet.name, cell.coordinate
                    ),
                    node_type_kind="formula",
                    sheet_name=sheet.name,
                    a1_range=A1Range.parse(sheet.name, cell.coordinate),
                    formula=cell.formula or "",
                    cached_value=cell.cached_value,
                    references=tuple(edge.reference for edge in parsed.references),
                    unresolved_references=parsed.unresolved,
                    member_coordinates=(cell.coordinate,),
                    is_cluster=False,
                )
            )
            continue
        min_row = min(cell.row for cell in cells)
        max_row = max(cell.row for cell in cells)
        min_col = min(cell.column for cell in cells)
        max_col = max(cell.column for cell in cells)
        a1 = _col_range_a1(min_col, max_col, min_row, max_row)
        entries.append(
            FormulaEntry(
                node_id=range_node_id(context.workbook_id, context.version, sheet.name, a1),
                node_type_kind="range",
                sheet_name=sheet.name,
                a1_range=A1Range.parse(sheet.name, a1),
                formula=cells[0].formula or "",
                cached_value=None,
                references=_expand_cluster_edges(
                    parsed.references,
                    sheet_name=sheet.name,
                    representative_row=cells[0].row,
                    group_min_row=min_row,
                    group_max_row=max_row,
                    workbook_id=context.workbook_id,
                    version=context.version,
                ),
                unresolved_references=parsed.unresolved,
                member_coordinates=tuple(cell.coordinate for cell in cells),
                is_cluster=True,
            )
        )
    entries.sort(key=lambda entry: entry.a1_range.a1)
    return tuple(entries)


def build_model(
    raw: RawWorkbook,
    *,
    workbook_id: str,
    version: int,
    acl_scope: Sequence[str] = (),
    config: RegionConfig | None = None,
) -> WorkbookModel:
    """Turn a :class:`RawWorkbook` into the canonical in-memory model."""
    settings = config or RegionConfig()
    known_sheets = frozenset(sheet.name for sheet in raw.sheets if not sheet.is_macro_sheet)
    macro_sheets = frozenset(raw.macro_sheet_names)
    tables = _build_tables(raw)
    named_ranges = _build_named_ranges(raw, workbook_id, version)
    defined_names = {
        name.name.lower(): NamedRangeInfo(
            name.name, name.sheet_name, A1Range.parse(name.sheet_name, name.a1)
        )
        for name in named_ranges
        if name.resolved and name.sheet_name and name.a1
    }
    sheet_max_row = {
        sheet.name: (max((cell.row for cell in sheet.cells.values()), default=1))
        for sheet in raw.sheets
    }

    sheet_max_col = {
        sheet.name: (max((cell.column for cell in sheet.cells.values()), default=1))
        for sheet in raw.sheets
    }
    sheet_order = tuple(sheet.name for sheet in raw.sheets)

    sheets: list[SheetModel] = []
    for sheet in raw.sheets:
        context = FormulaContext(
            workbook_id=workbook_id,
            version=version,
            sheet_name=sheet.name,
            known_sheets=known_sheets,
            macro_sheets=macro_sheets,
            sheet_max_row=sheet_max_row,
            tables=tables,
            defined_names=defined_names,
            sheet_order=sheet_order,
            sheet_max_col=sheet_max_col,
        )
        regions: tuple[Region, ...] = detect_regions(
            sheet, workbook_id=workbook_id, version=version, config=settings
        )
        formulas = () if sheet.is_macro_sheet else _build_formulas(sheet, context)
        sheets.append(
            SheetModel(
                name=sheet.name,
                visibility=sheet.visibility,
                node_id=sheet_node_id(workbook_id, version, sheet.name),
                a1_range=_used_range(sheet) if not sheet.is_macro_sheet else None,
                regions=regions,
                formulas=formulas,
                cells=sheet.cells,
                declared_dimension=sheet.declared_dimension,
                declared_dimension_flagged=sheet.declared_dimension_flagged,
                is_macro_sheet=sheet.is_macro_sheet,
                table_names=tuple(table.name for table in sheet.tables),
            )
        )
    return WorkbookModel(
        workbook_id=workbook_id,
        version=version,
        source_file=raw.source_file,
        source_sha256=raw.source_sha256,
        sheets=tuple(sheets),
        named_ranges=named_ranges,
        macro_sheet_names=raw.macro_sheet_names,
        has_vba=raw.has_vba,
        acl_scope=tuple(acl_scope),
        warnings=raw.warnings,
    )


__all__ = ["build_model", "normalize_formula"]
