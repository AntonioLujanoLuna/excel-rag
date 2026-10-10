"""The reader: what is loaded, what is refused, and what is only ever *noticed*."""

from __future__ import annotations

import pytest

from excel_rag.ingest import IngestError, read_workbook
from excel_rag.workbook.reader import merged_top_left


def test_reads_cells_merged_tables_and_formulas(build) -> None:
    raw = read_workbook(build.path("two_tables_one_sheet"))
    assert [sheet.name for sheet in raw.sheets] == ["Report"]
    sheet = raw.sheets[0]
    assert sheet.cells["A1"].value == "Quarterly Report"
    assert sheet.cells["A1"].merged_range == "A1:C1"
    assert sheet.cells["B4"].value == 100
    assert sheet.visibility == "visible"


def test_reads_table_objects(build) -> None:
    raw = read_workbook(build.path("table_object"))
    tables = raw.sheets[0].tables
    assert [table.name for table in tables] == ["SalesTable"]
    assert tables[0].ref == "A1:B4"
    assert tables[0].columns == ("Region", "Revenue")


def test_reads_defined_names(build) -> None:
    raw = read_workbook(build.path("named_range"))
    names = {defined.name: defined.attr_text for defined in raw.defined_names}
    assert names == {"GrowthRate": "Assumptions!$C$7"}


def test_formula_and_cached_value_are_kept_separate(build) -> None:
    raw = read_workbook(build.path("cached_value"))
    cell = raw.sheets[0].cells["C1"]
    assert cell.formula == "=A1+B1"
    assert cell.cached_value == 999
    assert cell.value is None  # a formula cell's own value is the formula, never a result


def test_visibility_is_carried(build) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Shown"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "x"
    path = build.directory / "visibility.xlsx"
    wb.save(path)
    raw = read_workbook(path)
    states = {sheet.name: sheet.visibility for sheet in raw.sheets}
    assert states == {"Shown": "visible", "Hidden": "hidden"}


def test_macro_and_vba_are_noticed_but_not_loaded(build) -> None:
    raw = read_workbook(build.path("macro_workbook"))
    assert raw.has_vba is True
    assert raw.macro_sheet_names == ("Macro1",)
    macro = next(sheet for sheet in raw.sheets if sheet.name == "Macro1")
    assert macro.is_macro_sheet is True
    assert macro.cells == {}  # a macro sheet's cells are never read


def test_hostile_declared_dimension_is_flagged_not_allocated(build) -> None:
    raw = read_workbook(build.path("hostile_dimension"))
    sheet = raw.sheets[0]
    assert sheet.declared_dimension == "A1:XFD1048576"
    assert sheet.declared_dimension_flagged is True
    assert len(sheet.cells) == 4  # only the cells that exist were built
    assert raw.warnings


def test_a_malformed_declared_dimension_is_ignored(build) -> None:
    import zipfile

    base = build.path("hostile_dimension")
    with zipfile.ZipFile(base) as archive:
        name = next(item for item in archive.namelist() if item.startswith("xl/worksheets/"))
    target = build.directory / "malformed_dimension.xlsx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(target, "w") as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == name:
                data = data.replace(b'ref="A1:XFD1048576"', b'ref="totally-not-a-range"')
            zout.writestr(item, data)
    raw = read_workbook(target)
    assert raw.sheets[0].declared_dimension_flagged is False


def test_corrupt_package_fails_loudly(build) -> None:
    with pytest.raises(IngestError, match="not a valid"):
        read_workbook(build.path("corrupt_file"))


def test_missing_file_fails_loudly(build) -> None:
    with pytest.raises(IngestError, match="no such workbook"):
        read_workbook(build.directory / "does_not_exist.xlsx")


def test_empty_sheet_is_read_as_empty(build) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    path = build.directory / "empty.xlsx"
    wb.save(path)
    raw = read_workbook(path)
    assert raw.sheets[0].cells == {}


@pytest.mark.parametrize(
    ("reference", "expected"),
    [("A1:B1", "A1"), ("$C$7:$D$9", "C7"), ("F18", "F18")],
)
def test_merged_top_left(reference: str, expected: str) -> None:
    assert merged_top_left(reference) == expected


def test_macro_workbook_via_helper(build) -> None:
    raw = read_workbook(build.path("macro_workbook"))
    assert any(sheet.is_macro_sheet for sheet in raw.sheets)


def test_source_metadata_is_recorded(build) -> None:
    raw = read_workbook(build.path("two_tables_one_sheet"))
    assert raw.source_file == "two_tables.xlsx"
    assert len(raw.source_sha256) == 64


_LAUGHS = (
    b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
    b'<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
)


@pytest.mark.parametrize("part", ["[Content_Types].xml", "xl/workbook.xml"])
def test_an_entity_declaring_part_is_refused_not_expanded(build, tmp_path, part) -> None:
    from fixtures.make_fixtures import _rewrite_zip

    def bomb(data: bytes) -> bytes:
        body = data.split(b"?>", 1)[1] if data.startswith(b"<?xml") else data
        return _LAUGHS + body

    path = _rewrite_zip(build.path("named_range"), tmp_path / "laughs.xlsx", edits={part: bomb})
    with pytest.raises(IngestError, match="XML entities"):
        read_workbook(path)


def test_a_part_that_is_not_xml_fails_loudly(build, tmp_path) -> None:
    from fixtures.make_fixtures import _rewrite_zip

    path = _rewrite_zip(
        build.path("named_range"),
        tmp_path / "broken.xlsx",
        edits={"[Content_Types].xml": b"<Types"},
    )
    with pytest.raises(IngestError, match="not XML"):
        read_workbook(path)
