"""Fixture workbooks for the context renderer and the workbook tools, built at test time."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "fixtures"))

import make_fixtures


@pytest.fixture
def fixtures(tmp_path: Path):  # type: ignore[no-untyped-def]
    """The shared fixture module, writing into this test's temporary directory."""

    class _Bound:
        def __getattr__(self, name: str) -> Path:
            builder = getattr(make_fixtures, name)
            return lambda **kwargs: builder(tmp_path, **kwargs)  # type: ignore[return-value]

    return _Bound()


@pytest.fixture
def typed_values(tmp_path: Path) -> Path:
    """Dates, percentages, booleans, a pipe and a long text: the values a grid must render."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Item", "Due", "Done", "Share", "Comment"])
    ws.append(["Launch", datetime(2026, 3, 31), True, 0.25, "go | no-go " + "x" * 200])
    ws.append(["Review", datetime(2026, 4, 2, 15, 30), False, 0.125, "ok"])
    for row in (2, 3):
        ws.cell(row=row, column=4).number_format = "0%"
    ws["H10"] = "stray note"
    path = tmp_path / "typed.xlsx"
    wb.save(path)
    return path
