"""Diffing two versions of a workbook: what changed, what it reads, and what reads it."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from excel_rag.cli import main
from excel_rag.workbook import load_workbook
from excel_rag.workbook.diff import ChangeKind, diff_files, diff_workbooks, format_diff


def _plan(directory: Path, name: str, *, growth: float, revenue_formula: str, extra: bool) -> Path:
    wb = openpyxl.Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"], inputs["B1"] = "Growth", growth
    inputs["A2"], inputs["B2"] = "Base", 100
    wb.defined_names["Growth"] = DefinedName("Growth", attr_text="Inputs!$B$1")
    model = wb.create_sheet("Model")
    model["A1"], model["B1"] = "Revenue", revenue_formula
    model["A2"], model["B2"] = "Label", "kept"
    if extra:
        wb.create_sheet("Scenarios")["A1"] = "new"
        model["A3"] = "added note"
    else:
        wb.create_sheet("Old")["A1"] = "gone"
    path = directory / name
    wb.save(path)
    return path


def _pair(tmp_path: Path) -> tuple[Path, Path]:
    before = _plan(
        tmp_path, "v1.xlsx", growth=0.05, revenue_formula="=Inputs!B2*(1+Growth)", extra=False
    )
    after = _plan(
        tmp_path, "v2.xlsx", growth=0.07, revenue_formula="=Inputs!B2*(1+Growth)*2", extra=True
    )
    return before, after


def test_an_identical_workbook_has_no_differences(tmp_path: Path) -> None:
    before, _ = _pair(tmp_path)
    diff = diff_files(before, before)
    assert diff.unchanged
    assert "No differences" in format_diff(diff)


def test_changes_are_classified_and_impact_is_traced(tmp_path: Path) -> None:
    diff = diff_files(*_pair(tmp_path))
    assert diff.sheets_added == ("Scenarios",)
    assert diff.sheets_removed == ("Old",)
    by_cell = {(change.sheet, change.coordinate): change for change in diff.changes}

    growth = by_cell[("Inputs", "B1")]
    assert growth.kind is ChangeKind.VALUE
    assert (growth.before, growth.after) == ("0.05", "0.07")
    assert growth.read_by == ("Model!B1",)

    revenue = by_cell[("Model", "B1")]
    assert revenue.kind is ChangeKind.FORMULA
    assert revenue.after_formula == "=Inputs!B2*(1+Growth)*2"

    assert by_cell[("Model", "A3")].kind is ChangeKind.ADDED
    assert ("Model", "B2") not in by_cell
    assert diff.counts == {"Inputs": {"value": 1}, "Model": {"added": 1, "formula": 1}}


def test_reads_added_and_removed_for_a_changed_formula(tmp_path: Path) -> None:
    before = _plan(tmp_path, "a.xlsx", growth=0.05, revenue_formula="=Inputs!B2", extra=False)
    after = _plan(tmp_path, "b.xlsx", growth=0.05, revenue_formula="=Inputs!B1", extra=False)
    (change,) = diff_files(before, after).changes
    assert change.reads_added == ("Inputs!B1",)
    assert change.reads_removed == ("Inputs!B2",)


def test_a_changed_saved_result_is_labelled_not_recomputed(build) -> None:
    model = load_workbook(build.path("cached_value"))
    sheet = model.sheets[0]
    formula_cell = next(cell for cell in sheet.cells.values() if cell.formula)
    altered_cells = dict(sheet.cells)
    from dataclasses import replace

    altered_cells[formula_cell.coordinate] = replace(formula_cell, cached_value=1)
    altered = replace(model, sheets=(replace(sheet, cells=altered_cells),))
    (change,) = diff_workbooks(model, altered).changes
    assert change.kind is ChangeKind.RESULT
    assert "formula unchanged" in format_diff(diff_workbooks(model, altered))


def test_a_value_turned_into_a_formula(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(tmp_path / "a.xlsx")
    wb.active["A1"] = "=1+0"
    wb.save(tmp_path / "b.xlsx")
    (change,) = diff_files(tmp_path / "a.xlsx", tmp_path / "b.xlsx").changes
    assert change.kind is ChangeKind.KIND


def test_type_changes_count_even_when_python_calls_them_equal(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = True
    wb.save(tmp_path / "a.xlsx")
    wb.active["A1"] = 1
    wb.save(tmp_path / "b.xlsx")
    (change,) = diff_files(tmp_path / "a.xlsx", tmp_path / "b.xlsx").changes
    assert change.kind is ChangeKind.VALUE


def test_named_range_changes(tmp_path: Path) -> None:
    before, after = _pair(tmp_path)
    wb = openpyxl.load_workbook(after)
    del wb.defined_names["Growth"]
    wb.defined_names["Growth"] = DefinedName("Growth", attr_text="Inputs!$B$2")
    wb.save(after)
    (named,) = diff_files(before, after).named_ranges
    assert (named.name, named.before, named.after) == ("Growth", "Inputs!B1", "Inputs!B2")


def test_the_listing_is_bounded_but_the_counts_are_not(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    for row in range(1, 21):
        wb.active[f"A{row}"] = row
    wb.save(tmp_path / "a.xlsx")
    for row in range(1, 21):
        wb.active[f"A{row}"] = row * 10
    wb.save(tmp_path / "b.xlsx")
    diff = diff_files(tmp_path / "a.xlsx", tmp_path / "b.xlsx", max_changes=5)
    assert len(diff.changes) == 5
    assert diff.total == 20
    assert diff.truncated
    assert "15 more change(s)" in format_diff(diff)


def test_cli_prints_text_and_json(tmp_path: Path, capsys) -> None:
    before, after = _pair(tmp_path)
    assert main(["diff", str(before), str(after)]) == 0
    assert "Inputs!B1" in capsys.readouterr().out
    assert main(["diff", str(before), str(after), "--json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["sheets_added"] == ["Scenarios"]
    assert payload["total"] == 3
