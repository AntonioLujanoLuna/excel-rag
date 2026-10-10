"""Structured references and whole-column names, end to end: what a formula reads is the column
(and the rows) it names, so ``dependents`` of a cell lists only the formulas that read it."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from excel_rag.context import WorkbookSession
from excel_rag.models import ReferenceKind
from excel_rag.workbook import load_workbook


def _sales_workbook(directory: Path) -> Path:
    """``Sales`` on Data!A1:D6: header, four data rows, a totals row; a calculated Amount column;
    a Calc sheet reading the table and a whole-column name."""
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Region", "Units", "Price", "Amount"])
    for region, units, price in [
        ("North", 3, 2.5),
        ("South", 5, 2.0),
        ("East", 2, 4.0),
        ("West", 7, 1.5),
    ]:
        data.append(
            [region, units, price, "=Sales[[#This Row],[Units]]*Sales[[#This Row],[Price]]"]
        )
    data.append(["Total", None, None, "=SUBTOTAL(109,Sales[Amount])"])
    table = Table(displayName="Sales", ref="A1:D6", totalsRowCount=1)
    data.add_table(table)
    calc = workbook.create_sheet("Calc")
    calc["A1"] = "=SUM(Sales[Amount])"
    calc["A2"] = "=COUNTA(Regions)"
    calc["A3"] = "=Sales[[#Totals],[Amount]]"
    workbook.defined_names["Regions"] = DefinedName("Regions", attr_text="Data!$A:$A")
    path = directory / "sales.xlsx"
    workbook.save(path)
    return path


def test_calculated_column_reads_its_own_row_band(tmp_path: Path) -> None:
    data = load_workbook(_sales_workbook(tmp_path)).sheets[0]
    cluster = next(entry for entry in data.formulas if entry.a1_range.a1 == "D2:D5")
    assert cluster.is_cluster
    assert {reference.a1_range for reference in cluster.references} == {"B2:B5", "C2:C5"}


def test_table_column_reference_is_the_columns_data_rows(tmp_path: Path) -> None:
    calc = load_workbook(_sales_workbook(tmp_path)).sheets[1]
    total = next(entry for entry in calc.formulas if entry.a1_range.a1 == "A1")
    assert [(r.a1_range, r.kind) for r in total.references] == [
        ("D2:D5", ReferenceKind.TABLE_COLUMN)
    ]
    whole_column = next(entry for entry in calc.formulas if entry.a1_range.a1 == "A2")
    assert [(r.a1_range, r.kind) for r in whole_column.references] == [
        ("A1:A6", ReferenceKind.NAMED_RANGE)
    ]
    totals = next(entry for entry in calc.formulas if entry.a1_range.a1 == "A3")
    assert [r.a1_range for r in totals.references] == ["D6"]


def test_dependents_list_only_the_formulas_reading_the_cell(tmp_path: Path) -> None:
    session = WorkbookSession.load(_sales_workbook(tmp_path))
    region = session.dependents("Data", "A3")
    assert "Calc!A2" in region  # COUNTA(Regions) reads column A
    assert "SUM(Sales[Amount])" not in region  # ... SUM(Sales[Amount]) does not
    units = session.dependents("Data", "B3")
    assert "Data!D2:D5" in units
    assert "SUM(Sales[Amount])" not in units
    amount = session.dependents("Data", "D3")
    assert "SUM(Sales[Amount])" in amount
    assert "Calc!A3" not in amount  # reads the totals row only
