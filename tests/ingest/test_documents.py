"""Document building: hierarchy, chunk text, bounded counts and no dangling edge targets."""

from __future__ import annotations

import pytest

from excel_rag.ingest import ingest_workbook
from excel_rag.ingest.build import normalize_formula
from excel_rag.ingest.canonical import workbook_node_id
from excel_rag.models import ChunkType, NodeType


def test_workbook_and_sheet_summaries_exist(build) -> None:
    ingested = build.ingest("two_tables_one_sheet")
    chunk_types = {chunk.chunk_type for chunk in ingested.chunks}
    assert {
        ChunkType.WORKBOOK,
        ChunkType.SHEET,
        ChunkType.TABLE,
        ChunkType.COLUMN,
        ChunkType.ROW_GROUP,
    } <= chunk_types


def test_structure_hierarchy_is_wired(build) -> None:
    ingested = ingest_workbook(build.path("merged_multilevel_header"), workbook_id="wb", version=1)
    nodes = {node.node_id: node for node in ingested.structure}
    workbook = nodes[workbook_node_id("wb", 1)]
    assert workbook.node_type is NodeType.WORKBOOK
    sheet_ids = [child for child in workbook.child_ids if nodes[child].node_type is NodeType.SHEET]
    assert len(sheet_ids) == 1
    sheet = nodes[sheet_ids[0]]
    assert any(nodes[child].node_type is NodeType.TABLE for child in sheet.child_ids)


def test_row_group_text_carries_its_column_headers(build) -> None:
    ingested = ingest_workbook(build.path("large_region"), workbook_id="wb", version=1)
    row_groups = [chunk for chunk in ingested.chunks if chunk.chunk_type is ChunkType.ROW_GROUP]
    assert row_groups
    for chunk in row_groups:
        assert "col_0" in chunk.content
        assert chunk.headers[:1] == ("col_0",)
        assert "Data" in chunk.title or "Data" in chunk.content


def test_region_text_names_sheet_and_a1_range(build) -> None:
    ingested = build.ingest("two_tables_one_sheet")
    tables = [chunk for chunk in ingested.chunks if chunk.chunk_type is ChunkType.TABLE]
    assert tables
    assert any("A3:C5" in chunk.content and "Report" in chunk.content for chunk in tables)


def test_no_cell_per_document_explosion(build) -> None:
    ingested = ingest_workbook(build.path("large_region"), workbook_id="wb", version=1)
    total = len(ingested.chunks) + len(ingested.structure)
    assert total < 100  # a 500x6 grid (3000 cells) produces dozens of documents
    cell_nodes = [node for node in ingested.structure if node.node_type is NodeType.CELL]
    assert len(cell_nodes) <= 6  # only the six header cells matter


def test_every_reference_target_has_a_structure_node(build) -> None:
    ingested = ingest_workbook(build.path("cross_sheet_formula"), workbook_id="wb", version=1)
    node_ids = {node.node_id for node in ingested.structure}
    for node in ingested.structure:
        for reference in node.references:
            assert reference.target_node_id in node_ids, reference.target_node_id


def test_formula_document_keeps_formula_and_cached_value_apart(build) -> None:
    ingested = ingest_workbook(build.path("cached_value"), workbook_id="wb", version=1)
    formulas = [node for node in ingested.structure if node.node_type is NodeType.FORMULA]
    assert len(formulas) == 1
    assert formulas[0].formula == "=A1+B1"
    assert formulas[0].cached_value == 999


def test_formula_cluster_collapses_repeated_pattern(build) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Series"
    ws["A1"], ws["B1"] = "x", "double"
    for row in range(2, 202):
        ws[f"A{row}"] = row
        ws[f"B{row}"] = f"=A{row}*2"
    path = build.directory / "cluster.xlsx"
    wb.save(path)
    ingested = ingest_workbook(path, workbook_id="wb", version=1)
    formula_chunks = [
        chunk for chunk in ingested.chunks if chunk.chunk_type is ChunkType.FORMULA_SUMMARY
    ]
    assert len(formula_chunks) == 1
    assert "repeats across 200 cells" in formula_chunks[0].content
    # the cluster's relative reference expands over the whole column, one edge
    formula_nodes = [
        node for node in ingested.structure if node.node_type is NodeType.RANGE and node.formula
    ]
    assert formula_nodes
    edges = formula_nodes[0].references
    assert len(edges) == 1
    assert edges[0].a1_range == "A2:A201"


def test_named_range_node_is_present_and_resolved(build) -> None:
    ingested = ingest_workbook(build.path("named_range"), workbook_id="wb", version=1)
    named = [node for node in ingested.structure if node.node_type is NodeType.NAMED_RANGE]
    assert len(named) == 1
    assert named[0].named_range == "GrowthRate"
    assert named[0].references[0].target_node_id in {node.node_id for node in ingested.structure}


def test_macro_sheet_is_documented_but_not_expanded(build) -> None:
    ingested = ingest_workbook(build.path("macro_workbook"), workbook_id="wb", version=1)
    sheets = [chunk for chunk in ingested.chunks if chunk.chunk_type is ChunkType.SHEET]
    assert any("macro sheet" in chunk.content.lower() for chunk in sheets)


def test_chunk_ids_are_unique_and_deterministic(build) -> None:
    first = ingest_workbook(build.path("two_tables_one_sheet"), workbook_id="wb", version=3)
    second = ingest_workbook(build.path("two_tables_one_sheet"), workbook_id="wb", version=3)
    ids = [chunk.id for chunk in first.chunks]
    assert len(ids) == len(set(ids))
    assert ids == [chunk.id for chunk in second.chunks]


def test_acl_scope_and_provenance_are_stamped(build) -> None:
    ingested = ingest_workbook(
        build.path("two_tables_one_sheet"), workbook_id="wb", version=1, acl_scope=("finance",)
    )
    assert all(chunk.acl_scope == ("finance",) for chunk in ingested.chunks)
    assert all(chunk.source_file == "two_tables.xlsx" for chunk in ingested.chunks)
    assert all(chunk.source_sha256 for chunk in ingested.chunks)


@pytest.mark.parametrize(
    ("left", "right", "same"),
    [
        ("=B2*$B$1", "=B3*$B$1", True),
        ("=B2+100", "=B3+200", False),
        ("=LOG10(A2)", "=LOG10(A3)", True),
        ('="Q1"&A2', '="Q2"&A3', False),
        ("=SUM(A2:C2)", "=SUM(A3:C3)", True),
        ("=$A2*2", "=$A3*2", True),
    ],
)
def test_cluster_patterns_rewrite_references_only(left: str, right: str, same: bool) -> None:
    assert (normalize_formula(left) == normalize_formula(right)) is same


def test_cluster_pattern_keeps_the_formula_text_around_references() -> None:
    assert normalize_formula('=IF(A2>0,"A2",LOG10(B2))') == '=IF(A#>0,"A2",LOG10(B#))'
