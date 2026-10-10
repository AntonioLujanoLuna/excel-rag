"""Array formulas, what-if data tables and sheet-scoped defined names, as Excel saves them."""

from __future__ import annotations

from excel_rag.models import NodeType, ReferenceKind, UnresolvedReason
from excel_rag.workbook import load_workbook
from excel_rag.workbook.reader import read_workbook


class TestArrayFormulas:
    def test_master_cell_carries_the_formula_text_not_an_object_repr(self, build) -> None:
        raw = read_workbook(build.path("array_formula"))
        master = raw.sheets[0].cells["B2"]
        assert master.formula == "=A2:A4*2"
        assert master.array_range == "B2:B4"
        assert master.cached_value == 2
        assert "object at 0x" not in (master.formula or "")

    def test_extent_cells_are_computed_cached_values_not_inputs(self, build) -> None:
        cells = read_workbook(build.path("array_formula")).sheets[0].cells
        for coordinate, saved in (("B3", 4), ("B4", 6)):
            cell = cells[coordinate]
            assert cell.value is None
            assert cell.cached_value == saved
            assert cell.array_master == "B2"
            assert cell.computed

    def test_array_formula_resolves_its_precedents(self, build) -> None:
        sheet = load_workbook(build.path("array_formula")).sheets[0]
        entry = next(entry for entry in sheet.formulas if entry.a1_range.a1 == "B2")
        assert not entry.is_cluster
        assert entry.array_range == "B2:B4"
        assert [reference.a1_range for reference in entry.references] == ["A2:A4"]
        assert entry.unresolved_references == ()

    def test_data_table_reads_its_input_cell_and_says_what_it_is(self, build) -> None:
        sheet = load_workbook(build.path("array_formula")).sheets[0]
        entry = next(entry for entry in sheet.formulas if entry.a1_range.a1 == "E2")
        assert entry.formula == "=TABLE(A2,)"
        assert entry.array_range == "E2:E3"
        assert [reference.a1_range for reference in entry.references] == ["A2"]
        assert {gap.reason for gap in entry.unresolved_references} == {
            UnresolvedReason.UNSUPPORTED_FUNCTION
        }
        assert sheet.cells["E3"].array_master == "E2"

    def test_reading_is_deterministic(self, build) -> None:
        path = build.path("array_formula")
        first = load_workbook(path).sheets[0].formulas
        second = load_workbook(path).sheets[0].formulas
        assert first == second

    def test_formula_summary_states_the_saved_extent(self, build) -> None:
        ingested = build.ingest("array_formula")
        summary = next(chunk for chunk in ingested.chunks if "A2:A4*2" in chunk.content)
        assert "last covered B2:B4" in summary.content

    def test_rendering_marks_extent_values_as_formula_results(self, build) -> None:
        from excel_rag.context import render_workbook

        text = render_workbook(build.path("array_formula")).text
        assert "(array over B2:B4)" in text
        row3 = next(line for line in text.splitlines() if line.startswith("| 3 "))
        assert "4 ƒ" in row3


class TestSheetScopedNames:
    def test_reader_keeps_sheet_scoped_names(self, build) -> None:
        raw = read_workbook(build.path("sheet_scoped_names"))
        names = {(defined.scope_sheet, defined.name) for defined in raw.defined_names}
        assert names == {(None, "Tax"), ("Jan", "Rate"), ("Feb", "Rate"), ("Feb", "Tax")}

    def test_same_local_name_on_two_sheets_is_two_nodes(self, build) -> None:
        model = load_workbook(build.path("sheet_scoped_names"), workbook_id="wb", version=1)
        ids = {name.label: name.node_id for name in model.named_ranges}
        assert ids["Jan!Rate"] == "wb:v1:named_range:jan!rate"
        assert ids["Feb!Rate"] == "wb:v1:named_range:feb!rate"
        assert ids["Tax"] == "wb:v1:named_range:tax"
        assert len(set(ids.values())) == 4

    def _targets(self, model, sheet: str) -> list[tuple[str, str, str]]:
        entry = next(s for s in model.sheets if s.name == sheet).formulas[0]
        assert entry.unresolved_references == ()
        return [
            (reference.target_node_id, reference.sheet_name, reference.a1_range)
            for reference in entry.references
            if reference.kind is ReferenceKind.NAMED_RANGE
        ]

    def test_a_local_name_resolves_on_its_own_sheet(self, build) -> None:
        model = load_workbook(build.path("sheet_scoped_names"), workbook_id="wb", version=1)
        assert self._targets(model, "Jan") == [("wb:v1:named_range:jan!rate", "Jan", "C1")]

    def test_a_local_name_shadows_the_global_one(self, build) -> None:
        model = load_workbook(build.path("sheet_scoped_names"), workbook_id="wb", version=1)
        assert self._targets(model, "Feb") == [("wb:v1:named_range:feb!tax", "Feb", "C2")]

    def test_qualified_and_global_names_from_another_sheet(self, build) -> None:
        model = load_workbook(build.path("sheet_scoped_names"), workbook_id="wb", version=1)
        assert self._targets(model, "Summary") == [
            ("wb:v1:named_range:jan!rate", "Jan", "C1"),
            ("wb:v1:named_range:feb!rate", "Feb", "C1"),
            ("wb:v1:named_range:tax", "Jan", "C2"),
        ]

    def test_every_name_is_a_structure_node(self, build) -> None:
        ingested = build.ingest("sheet_scoped_names")
        named = sorted(
            node.node_id for node in ingested.structure if node.node_type is NodeType.NAMED_RANGE
        )
        assert named == [
            "wb:v1:named_range:feb!rate",
            "wb:v1:named_range:feb!tax",
            "wb:v1:named_range:jan!rate",
            "wb:v1:named_range:tax",
        ]


def test_a_hostile_array_extent_costs_a_pass_over_the_cells(tmp_path) -> None:
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = ArrayFormula("A1:XFD1048576", "=1")
    ws["B2"] = 5
    path = tmp_path / "hostile_array.xlsx"
    wb.save(path)
    cells = read_workbook(path).sheets[0].cells
    assert cells["B2"].array_master == "A1"
    assert cells["B2"].cached_value == 5
    assert set(cells) == {"A1", "B2"}
